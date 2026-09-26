# -*- coding: utf-8 -*-
"""Export a Shanghai high-school English exam to Word (.docx)."""

from __future__ import annotations

import json
import re
from pathlib import Path
from typing import Any

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

ASCII = set(range(32, 127))
LISTENING_GROUP = re.compile(
    r"Questions?\s+\d+\s*(?:through|to|-|–|—)\s*\d+\s+are\s+based\s+on\s+the\s+following"
    r"\s+(?:longer\s+)?(?:passage|conversation)\.?",
    re.I,
)
CHOICE_LINE = re.compile(r"^\s*([A-D])[\.、．]\s*(.*)$")
NUMBERED = re.compile(r"^\s*(\d+)[\.、．]\s*(.*)$")
BULLET = re.compile(r"^\s*[•·\-]\s+(.*)$")


def split_runs(text: str) -> list[tuple[str, bool]]:
    """Split into (token, is_ascii) runs so Word can mix Times / 宋体."""
    parts: list[tuple[str, bool]] = []
    for tok in re.findall(r"[ -~]+|[^\x20-\x7E]+", text or ""):
        is_ascii = bool(tok) and all(ord(ch) in ASCII for ch in tok)
        parts.append((tok, is_ascii))
    return parts or [("", True)]


def set_run_font(run, *, en: str, zh: str, size: float, bold: bool = False, italic: bool = False, underline: bool = False) -> None:
    run.bold = bold
    run.italic = italic
    run.underline = underline
    run.font.size = Pt(size)
    run.font.name = en
    r_pr = run._element.get_or_add_rPr()
    r_fonts = r_pr.get_or_add_rFonts()
    r_fonts.set(qn("w:ascii"), en)
    r_fonts.set(qn("w:hAnsi"), en)
    r_fonts.set(qn("w:eastAsia"), zh)
    r_fonts.set(qn("w:cs"), en)


def add_mixed_runs(
    paragraph,
    text: str,
    *,
    en: str,
    zh: str,
    size: float,
    bold: bool = False,
    italic: bool = False,
    underline: bool = False,
) -> None:
    for token, is_ascii in split_runs(text):
        run = paragraph.add_run(token)
        set_run_font(run, en=en, zh=zh, size=size, bold=bold, italic=italic, underline=underline)
        if not is_ascii:
            run.font.name = zh


def set_paragraph_format(paragraph, *, line_spacing: float, after_pt: float = 6, first_line_cm: float = 0) -> None:
    fmt = paragraph.paragraph_format
    fmt.line_spacing = line_spacing
    fmt.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    fmt.space_after = Pt(after_pt)
    fmt.space_before = Pt(0)
    if first_line_cm:
        fmt.first_line_indent = Cm(first_line_cm)


def load_exam(path: str | Path) -> dict[str, Any]:
    data = json.loads(Path(path).read_text(encoding="utf-8"))
    if "exam" not in data or "sections" not in data:
        raise ValueError("试卷 JSON 需要包含 exam 和 sections 字段。")
    if not isinstance(data["sections"], list) or not data["sections"]:
        raise ValueError("sections 不能为空。")
    return data


def _style(data: dict[str, Any]) -> dict[str, Any]:
    style = data.get("style") or {}
    return {
        "en": style.get("font_en") or "Times New Roman",
        "zh": style.get("font_zh") or "宋体",
        "size": float(style.get("size_pt") or 10.5),
        "line": float(style.get("line_spacing") or 1.2),
        "indent": float(style.get("first_line_indent_cm") or 0.74),
    }


def _add_para(
    doc: Document,
    text: str,
    style: dict[str, Any],
    *,
    bold: bool = False,
    italic: bool = False,
    center: bool = False,
    justify: bool = False,
    indent: bool = False,
    size: float | None = None,
    after_pt: float = 6,
    before_pt: float = 0,
) -> None:
    para = doc.add_paragraph()
    if center:
        para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    elif justify:
        para.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    set_paragraph_format(
        para,
        line_spacing=style["line"],
        after_pt=after_pt,
        first_line_cm=style["indent"] if indent else 0,
    )
    if before_pt:
        para.paragraph_format.space_before = Pt(before_pt)
    add_mixed_runs(
        para,
        text,
        en=style["en"],
        zh=style["zh"],
        size=size or style["size"],
        bold=bold,
        italic=italic,
    )


def _render_content_line(doc: Document, line: str, section_type: str, style: dict[str, Any]) -> None:
    stripped = line.rstrip()
    if not stripped.strip():
        return

    if section_type == "listening" and LISTENING_GROUP.search(stripped):
        _add_para(doc, stripped.strip(), style, bold=True, italic=True, after_pt=2)
        return

    bullet = BULLET.match(stripped)
    if bullet:
        para = doc.add_paragraph(style="List Bullet")
        set_paragraph_format(para, line_spacing=style["line"], after_pt=2)
        add_mixed_runs(para, bullet.group(1), en=style["en"], zh=style["zh"], size=style["size"])
        return

    choice = CHOICE_LINE.match(stripped)
    if choice:
        _add_para(doc, f"{choice.group(1)}. {choice.group(2)}", style, after_pt=2)
        return

    numbered = NUMBERED.match(stripped)
    indent = section_type not in {"translation", "listening"}
    if numbered:
        _add_para(
            doc,
            f"{numbered.group(1)}. {numbered.group(2)}",
            style,
            justify=True,
            indent=False,
            after_pt=4,
        )
        return

    _add_para(
        doc,
        stripped.strip(),
        style,
        justify=True,
        indent=indent and section_type in {"reading", "cloze", "summary", "writing"},
        after_pt=4,
    )


def _render_word_bank(doc: Document, words: list[str], style: dict[str, Any]) -> None:
    if not words:
        return
    table = doc.add_table(rows=1, cols=min(5, len(words)))
    table.autofit = True
    row = table.rows[0]
    for i, cell in enumerate(row.cells):
        if i >= len(words):
            break
        cell.text = ""
        para = cell.paragraphs[0]
        para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        add_mixed_runs(para, words[i], en=style["en"], zh=style["zh"], size=style["size"])
    remaining = words[len(row.cells) :]
    while remaining:
        row = table.add_row()
        chunk, remaining = remaining[: len(row.cells)], remaining[len(row.cells) :]
        for i, cell in enumerate(row.cells):
            cell.text = ""
            para = cell.paragraphs[0]
            para.alignment = WD_ALIGN_PARAGRAPH.CENTER
            if i < len(chunk):
                add_mixed_runs(para, chunk[i], en=style["en"], zh=style["zh"], size=style["size"])
    doc.add_paragraph()


def _render_section(doc: Document, section: dict[str, Any], style: dict[str, Any], last_major: str | None) -> str | None:
    major = str(section.get("major") or "").strip()
    major_title = str(section.get("major_title") or section.get("title") or "").strip()
    major_key = f"{major}|{major_title}" if major else None

    if major and major_key != last_major:
        _add_para(
            doc,
            f"{major}. {major_title}".strip(". "),
            style,
            bold=True,
            before_pt=10,
            after_pt=4,
        )
    elif not major:
        title = section.get("title") or "未命名大题"
        _add_para(doc, str(title), style, bold=True, before_pt=10, after_pt=4)

    label = section.get("section_label") or section.get("sectionLabel")
    if label:
        _add_para(doc, str(label), style, bold=True, after_pt=2)

    instruction = section.get("instruction")
    if instruction:
        _add_para(doc, str(instruction), style, bold=True, justify=True, after_pt=4)

    part_label = section.get("part_label") or section.get("reading_part_label")
    if part_label:
        reading = str(section.get("type") or "") in {"reading", "section_c"}
        _add_para(doc, str(part_label), style, bold=True, center=reading, after_pt=4)

    section_type = str(section.get("type") or "custom")
    word_bank = section.get("word_bank") or section.get("wordBank")
    if word_bank:
        _render_word_bank(doc, [str(w) for w in word_bank], style)

    items = section.get("items")
    if items:
        for item in items:
            if isinstance(item, str):
                _render_content_line(doc, item, section_type, style)
                continue
            no = item.get("no") or item.get("number")
            text = item.get("text") or item.get("chinese") or ""
            score = item.get("score")
            extra = f"（{score}分）" if score not in (None, "") else ""
            prefix = f"{no}. " if no not in (None, "") else ""
            _add_para(doc, f"{prefix}{text}{extra}", style, justify=True, after_pt=6)
            hints = item.get("hints") or item.get("words")
            if hints:
                _add_para(doc, f"（{hints}）", style, italic=True, after_pt=4)
        return major_key

    content = section.get("content") or ""
    for line in str(content).splitlines():
        _render_content_line(doc, line, section_type, style)
    return major_key


def _build_paper(doc: Document, data: dict[str, Any]) -> None:
    exam = data.get("exam") or {}
    style = _style(data)
    title = exam.get("title") or "试卷"
    _add_para(doc, str(title), style, bold=True, center=True, size=16, after_pt=8)

    score = exam.get("total_score") if "total_score" in exam else exam.get("totalScore")
    duration = exam.get("duration") or ""
    meta = f"（满分：{score or 0} 分   时间：{duration}）"
    year_month = exam.get("year_month") or exam.get("yearMonth")
    if year_month:
        para = doc.add_paragraph()
        para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_paragraph_format(para, line_spacing=style["line"], after_pt=12)
        add_mixed_runs(para, meta, en=style["en"], zh=style["zh"], size=style["size"])
        run = para.add_run(f"    {year_month}")
        set_run_font(run, en=style["en"], zh=style["zh"], size=style["size"])
    else:
        _add_para(doc, meta, style, center=True, after_pt=12)

    last_major = None
    for section in data.get("sections") or []:
        last_major = _render_section(doc, section, style, last_major)


def _build_answers(doc: Document, data: dict[str, Any]) -> None:
    exam = data.get("exam") or {}
    style = _style(data)
    title = f"{exam.get('title') or '试卷'} 答案"
    _add_para(doc, title, style, bold=True, center=True, size=16, after_pt=12)

    answers = data.get("answers") or []
    if not answers:
        _add_para(doc, "（尚未填写答案）", style, italic=True)
    for block in answers:
        heading = block.get("title") or block.get("section_heading") or ""
        if heading:
            _add_para(doc, str(heading), style, bold=True, before_pt=8, after_pt=4)
        text = block.get("text") or ""
        for line in str(text).splitlines() or [""]:
            if line.strip():
                _add_para(doc, line, style, justify=True, after_pt=4)
            else:
                _add_para(doc, "（尚未填写答案）", style, italic=True)

    scripts = data.get("listening_scripts") or data.get("listeningScripts") or []
    if scripts:
        _add_para(doc, "听力原文", style, bold=True, before_pt=12, after_pt=6)
        for block in scripts:
            label = block.get("label") or block.get("section_label") or ""
            if label:
                _add_para(doc, str(label), style, bold=True, after_pt=2)
            for line in str(block.get("text") or "").splitlines():
                if line.strip():
                    _add_para(doc, line, style, justify=True, indent=True, after_pt=4)


def export_exam_docx(data: dict[str, Any], output_path: str, answers: bool = False) -> str:
    doc = Document()
    for section in doc.sections:
        section.top_margin = Cm(2.54)
        section.bottom_margin = Cm(2.54)
        section.left_margin = Cm(2.54)
        section.right_margin = Cm(2.54)

    if answers:
        _build_answers(doc, data)
    else:
        _build_paper(doc, data)

    Path(output_path).parent.mkdir(parents=True, exist_ok=True)
    doc.save(output_path)
    return output_path


EXAMPLE_EXAM = {
    "exam": {
        "title": "2026学年第一学期高一年级英语期中考试",
        "total_score": 150,
        "duration": "120分钟",
        "year_month": "2026.11",
    },
    "style": {
        "font_en": "Times New Roman",
        "font_zh": "宋体",
        "size_pt": 10.5,
        "line_spacing": 1.2,
    },
    "sections": [
        {
            "major": "I",
            "major_title": "Listening Comprehension",
            "section_label": "Section A",
            "type": "listening",
            "instruction": "Directions: In Section A, you will hear ten short conversations between two speakers. At the end of each conversation, a question will be asked about what was said. Both the conversation and the question will be spoken only once. After you hear a conversation and the question about it, read the four possible answers on your paper, and decide which one is the best answer to the question you have heard.",
            "content": "1. A. In a library.\nB. In a bookstore.\nC. In a classroom.\nD. In a museum.\n2. A. Teacher and student.\nB. Doctor and patient.\nC. Waiter and customer.\nD. Driver and passenger.",
        },
        {
            "major": "I",
            "major_title": "Listening Comprehension",
            "section_label": "Section B",
            "type": "listening",
            "instruction": "Directions: In Section B, you will hear two passages and one longer conversation. After each, you will be asked several questions. The passages and the conversation will be read twice.",
            "content": "Questions 11 through 13 are based on the following passage.\n11. A. A school trip.\nB. A science project.\nC. A sports meeting.\nD. A music festival.",
        },
        {
            "major": "II",
            "major_title": "Grammar and Vocabulary",
            "section_label": "Section A",
            "type": "grammar",
            "instruction": "Directions: After reading the passage below, fill in the blanks to make the passage coherent and grammatically correct. For the blanks with a given word, fill in each blank with the proper form of the given word; for the other blanks, use one word that best fits each blank.",
            "content": "(A)\nWhen Tom first arrived in the city, he found (21)______ hard to make new friends. He kept telling himself that things (22)______ (get) better if he stayed patient.",
        },
        {
            "major": "II",
            "major_title": "Grammar and Vocabulary",
            "section_label": "Section B",
            "type": "vocabulary",
            "instruction": "Directions: Fill in each blank with a proper word chosen from the box. Each word can be used only once. Note that there is one word more than you need.",
            "word_bank": ["concern", "effective", "gradually", "maintain", "obvious", "range"],
            "content": "41. Regular exercise is an ______ way to reduce stress.\n42. The park offers a wide ______ of outdoor activities.",
        },
        {
            "major": "III",
            "major_title": "Reading Comprehension",
            "section_label": "Section A",
            "type": "cloze",
            "instruction": "Directions: For each blank in the following passage there are four words or phrases marked A, B, C and D. Fill in each blank with the word or phrase that best fits the context.",
            "content": "Learning a new language takes time, but small daily habits can (51) ______ a big difference.\n51. A. make  B. take  C. keep  D. find",
        },
        {
            "major": "III",
            "major_title": "Reading Comprehension",
            "section_label": "Section B",
            "type": "reading",
            "reading_part_label": "(A)",
            "instruction": "Directions: Read the following passages. Each passage is followed by several questions. For each of them there are four choices marked A, B, C and D. Choose the one that fits best according to the information given in the passage you have just read.",
            "content": "Many students think revision means reading notes again and again. Research shows that testing yourself is often more useful.\n65. What is the main idea of the passage?\nA. Reading notes is useless.\nB. Self-testing helps revision.\nC. Students dislike exams.\nD. Research is always right.",
        },
        {
            "major": "IV",
            "major_title": "Translation",
            "type": "translation",
            "instruction": "Directions: Translate the following sentences into English, using the words given in the brackets.",
            "items": [
                {"no": 72, "text": "只要你坚持练习，口语一定会有进步。", "hints": "as long as", "score": 4},
                {"no": 73, "text": "令我们惊讶的是，他最终接受了这个挑战。", "hints": "to one's surprise", "score": 4},
            ],
        },
        {
            "major": "V",
            "major_title": "Guided Writing",
            "type": "writing",
            "instruction": "Directions: Write an English composition in 120–150 words according to the instructions given below in Chinese.",
            "content": "学校将举办英语读书周。请你写一封邮件给外教 Mr. Green，内容包括：\n- 活动时间与地点\n- 需要他做的事\n- 你的期待",
        },
    ],
    "answers": [
        {"title": "I. Listening Comprehension", "text": "1-5 B C A D B\n11-13 A C B"},
        {"title": "II. Grammar and Vocabulary", "text": "21. it  22. would get\n41. effective  42. range"},
        {"title": "III. Reading Comprehension", "text": "51. A\n65. B"},
        {"title": "IV. Translation", "text": "72. As long as you keep practising, your spoken English is bound to improve.\n73. To our surprise, he finally took on the challenge."},
    ],
    "listening_scripts": [
        {
            "label": "Section A",
            "text": "1. W: Excuse me, where can I find the history books?\nM: They are on the second floor, next to the magazines.\nQ: Where does the conversation most probably take place?",
        }
    ],
}
