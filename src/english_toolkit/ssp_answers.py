# -*- coding: utf-8 -*-
"""Parse SSP 参考答案 PDF and export a teacher answer Word."""

from __future__ import annotations

import re
from pathlib import Path
from typing import Any

import pymupdf
from docx import Document
from docx.shared import Cm

from english_toolkit.exam import set_paragraph_format
from english_toolkit.ssp import _ascii_quotes, _norm_title, _para, _prepare_export_articles, _section_code

RANGE_KEY_RE = re.compile(r"^(\d+)\s*[-–—]\s*(\d+)\s+([A-Ka-k]+)$")
BARE_KEY_RE = re.compile(r"^([A-F]{3,8})$")
SECTION_HEAD_RE = re.compile(r"^([IVX]+)\.\s*(.*)$")
NUMBERED_ANS_RE = re.compile(r"^(\d+)\.\s+(.+)$")
CIRCLED_ANS_RE = re.compile(r"^([①②③④⑤⑥⑦⑧⑨⑩])\s*(.+)$")
PAGE_NOISE_RE = re.compile(r"^--\s*\d+\s+of\s+\d+\s*--$")
STOP_TITLES = {
    "sentence structures",
    "summary writing",
    "learn to write",
    "a possible answer:",
}


def parse_answer_pdf(path: str | Path) -> dict[str, Any]:
    doc = pymupdf.open(path)
    try:
        text = "\n".join(page.get_text() for page in doc)
    finally:
        doc.close()
    return parse_answer_text(text)


def parse_answer_text(text: str) -> dict[str, Any]:
    articles: list[dict[str, Any]] = []
    current: dict[str, Any] | None = None
    section = ""
    in_cloze = False
    stopped = False

    def bucket() -> dict[str, Any]:
        nonlocal current, section
        if current is None:
            start_article("未命名")
        key = "Cloze" if in_cloze else (section or "I")
        return current["sections"].setdefault(key, {"mcq": "", "lines": []})

    def start_article(title: str, *, cloze: bool = False) -> None:
        nonlocal current, section, in_cloze
        current = {
            "title": re.sub(r"\s*\([AB]\d\)\s*$", "", _ascii_quotes(title)).strip(),
            "sections": {},
            "cloze": cloze,
        }
        articles.append(current)
        section = "Cloze" if cloze else ""
        in_cloze = cloze

    def add_mcq(letters: str) -> None:
        data = bucket()
        data["mcq"] += letters.upper()

    for raw in text.splitlines():
        line = raw.strip()
        if not line or PAGE_NOISE_RE.match(line):
            continue
        lower = line.lower()
        if re.match(r"^Cloze$", line, re.I):
            stopped = False
            in_cloze = True
            section = "Cloze"
            current = None
            continue
        if lower in STOP_TITLES or lower.startswith("方向"):
            stopped = True
            continue
        if stopped:
            continue
        if line in {"B1", "B2", "B3", "B4"} or line.startswith("高中进阶版") or re.fullmatch(r"[AB]\d", line):
            continue
        if re.match(r"^（\d{4}", line):
            continue

        head = SECTION_HEAD_RE.match(line)
        if head and current is not None and not in_cloze:
            section = head.group(1)
            current["sections"].setdefault(section, {"mcq": "", "lines": []})
            after = head.group(2).strip()
            ranged = RANGE_KEY_RE.match(after)
            if ranged:
                add_mcq(ranged.group(3))
            elif BARE_KEY_RE.match(after):
                add_mcq(after)
            continue

        ranged = RANGE_KEY_RE.match(line)
        if ranged and current is not None:
            add_mcq(ranged.group(3))
            continue
        if current is not None and BARE_KEY_RE.match(line) and (section in {"III", "Cloze"} or in_cloze):
            add_mcq(line)
            continue

        if in_cloze and current is None and re.search(r"[A-Za-z]{4,}", line) and not ranged:
            start_article(line, cloze=True)
            continue

        if (
            current is not None
            and not in_cloze
            and re.match(r"^[A-Z“\"]", line)
            and len(line) >= 16
            and not NUMBERED_ANS_RE.match(line)
            and not CIRCLED_ANS_RE.match(line)
        ):
            start_article(line)
            continue
        if current is None and re.match(r"^[A-Z]", line) and len(line) >= 16:
            start_article(line)
            continue

        if current is None:
            continue
        if NUMBERED_ANS_RE.match(line) or CIRCLED_ANS_RE.match(line) or line.startswith(("①", "②", "③", "④")):
            bucket()["lines"].append(line)
            continue
        if section and bucket()["lines"] and not NUMBERED_ANS_RE.match(line):
            bucket()["lines"][-1] += " " + line

    return {"articles": articles, "text": text}


def _match_answer_article(title: str, parsed: dict[str, Any]) -> dict[str, Any] | None:
    want = _norm_title(title)
    best = None
    best_score = 0
    for article in parsed.get("articles") or []:
        have = _norm_title(article.get("title") or "")
        if not have:
            continue
        if want == have or want in have or have in want:
            return article
        tokens = [t for t in re.findall(r"[A-Za-z]{4,}", want) if t not in {"what", "like", "live", "want", "from"}]
        score = sum(1 for t in tokens if t in have)
        if score > best_score and score >= 2:
            best = article
            best_score = score
    return best


def apply_answer_key(mcqs: list[dict[str, Any]], parsed: dict[str, Any]) -> list[dict[str, Any]]:
    for item in mcqs:
        article = _match_answer_article(item.get("article_title") or "", parsed)
        if article is None:
            continue
        section = "Cloze" if item.get("is_cloze") else (item.get("section") or "")
        blob = ""
        if item.get("is_cloze"):
            blob = (article.get("sections") or {}).get("Cloze", {}).get("mcq") or ""
            if not blob:
                blob = "".join(sec.get("mcq") or "" for sec in (article.get("sections") or {}).values())
        else:
            blob = (article.get("sections") or {}).get(section, {}).get("mcq") or ""
        orig = str(item.get("orig_no") or "")
        if orig.isdigit() and blob:
            idx = int(orig) - 1
            if 0 <= idx < len(blob):
                item["answer"] = blob[idx].upper()
    return mcqs


def compact_mcq_answers(mcqs: list[dict[str, Any]]) -> str:
    return "".join(str(item.get("answer") or "?") for item in mcqs)


def export_reading_answers_docx(
    issue: dict[str, Any],
    article_ids: list[str],
    output_path: str | Path,
    *,
    answer_pdf: str | Path | None = None,
    include_vocab_exercises: bool = True,
    include_questions: bool = True,
) -> Path:
    lookup = {a["id"]: a for a in issue.get("articles") or []}
    ordered = [lookup[i] for i in article_ids if i in lookup]
    if not ordered:
        raise ValueError("没有选中的文章。")
    prepared, mcqs = _prepare_export_articles(
        ordered,
        include_vocab_exercises=include_vocab_exercises,
        include_questions=include_questions,
    )
    parsed: dict[str, Any] = {"articles": []}
    if answer_pdf:
        parsed = parse_answer_pdf(answer_pdf)
        apply_answer_key(mcqs, parsed)

    doc = Document()
    for section in doc.sections:
        section.page_width = Cm(21.0)
        section.page_height = Cm(29.7)
        section.top_margin = Cm(2.54)
        section.bottom_margin = Cm(2.54)
        section.left_margin = Cm(2.54)
        section.right_margin = Cm(2.54)
    style = {
        "en": "Times New Roman",
        "zh": "宋体",
        "size": 10.5,
        "line": 1.2,
        "indent": 0.74,
        "after": 0,
    }
    issue_no = issue.get("issue") or ""
    _para(doc, "高二英语阅读答案", style, bold=True, center=True, size=16, after=6)
    sub = "Shanghai Students' Post 高中进阶版"
    if issue_no:
        sub += f"  第{issue_no}期"
    _para(doc, sub, style, italic=True, center=True, after=8)

    if mcqs:
        _para(doc, "选择题（与答题卡题号一致）", style, bold=True, after=4)
        keys = compact_mcq_answers(mcqs)
        _para(doc, f"共 {len(mcqs)} 题：{keys}", style, after=6)
        chunks: list[str] = []
        start = 1
        last_title = ""
        last_block = ""
        buf: list[str] = []

        def flush() -> None:
            nonlocal start, buf
            if not buf:
                return
            end = start + len(buf) - 1
            rng = f"{start}" if start == end else f"{start}-{end}"
            _para(doc, f"{rng}  {''.join(buf)}    {last_title}  {last_block}".strip(), style, after=1)
            start = end + 1
            buf = []

        for item in mcqs:
            title = item.get("article_title") or ""
            block = item.get("block_title") or ""
            if buf and (title != last_title or block != last_block):
                flush()
            last_title = title
            last_block = block
            buf.append(str(item.get("answer") or "?"))
        flush()
        extra = doc.add_paragraph()
        set_paragraph_format(extra, line_spacing=style["line"], after_pt=8)

    for idx, article in enumerate(prepared):
        _para(doc, f"Passage {idx + 1}  {article.get('title') or ''}", style, bold=True, after=4)
        parsed_art = _match_answer_article(article.get("title") or "", parsed)
        for block in article.get("exercises") or []:
            title = block.get("title") or ""
            if title:
                _para(doc, title, style, bold=True, after=2)
            option_qs = [q for q in (block.get("questions") or []) if q.get("options")]
            if option_qs:
                bits = []
                for q in option_qs:
                    ans = next((m.get("answer") for m in mcqs if m.get("no") == q.get("no")), "") or "?"
                    bits.append(f"{q.get('no')}{ans}")
                _para(doc, "  ".join(bits), style, after=2)
            section = _section_code(title)
            lines = []
            if parsed_art and section:
                lines = list((parsed_art.get("sections") or {}).get(section, {}).get("lines") or [])
            if not option_qs:
                six = ""
                if parsed_art and section:
                    six = (parsed_art.get("sections") or {}).get(section, {}).get("mcq") or ""
                if six and re.fullmatch(r"[A-F]+", six):
                    _para(doc, six, style, after=2)
            for line in lines:
                _para(doc, line, style, after=1)
        extra = doc.add_paragraph()
        set_paragraph_format(extra, line_spacing=style["line"], after_pt=6)

    out = Path(output_path)
    out.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(out))
    return out
