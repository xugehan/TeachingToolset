# -*- coding: utf-8 -*-
"""Parse Shanghai Students' Post PDFs and export reading worksheets."""

from __future__ import annotations

import re
from dataclasses import asdict, dataclass, field
from io import BytesIO
from pathlib import Path
from typing import Any, Iterable

import pymupdf
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm

from english_toolkit.exam import add_mixed_runs, set_paragraph_format

META_RE = re.compile(
    r"语篇类型[:：]\s*(?P<genre>\S+)\s*[／/]\s*词数[:：]\s*(?P<words>\d+)\s*"
    r"[／/]\s*主题[:：]\s*(?P<theme>.+?)\s*[／/]\s*难度[:：]\s*(?P<level>\d+)\s*级"
)
PAGE_CODE_RE = re.compile(r"^[AB][1-8]$")
APOSTROPHE_RE = re.compile(r"\b([A-Za-z]+) s\b")
QUOTE_MAP = str.maketrans(
    {
        "‘": "'",
        "’": "'",
        "“": '"',
        "”": '"',
        "「": '"',
        "」": '"',
        "『": '"',
        "』": '"',
        "＇": "'",
        "＂": '"',
    }
)


def _ascii_quotes(text: str) -> str:
    return text.translate(QUOTE_MAP)
CHOICE_RE = re.compile(r"^([A-F])[\.、．]\s*(.*)$")
LABELED_WORD_RE = re.compile(r"^([A-K])[\.、．]\s*([A-Za-z][A-Za-z'-]*)$")
NUMBERED_RE = re.compile(r"^(\d+)[\.、．]\s*(.*)$")
SECTION_RE = re.compile(r"^([IVX]+)\.\s+")
ARTICLE_HEADER_RE = re.compile(r"^A(\d)\s+(.+)$")
CIRCLED_RE = re.compile(r"[①②③④⑤⑥⑦⑧⑨⑩]")
CIRCLED_ONLY_RE = re.compile(r"^[①②③④⑤⑥⑦⑧⑨⑩]$")
ITALIC_MARK = re.compile(r"\*[^*]+\*")
CONTENT_WIDTH_CM = 15.92
WORD_BANK_LINE_RE = re.compile(r"^[A-Za-z][A-Za-z'-]*(?:\s+[A-Za-z][A-Za-z'-]*){1,7}$")


def _strip_markup(text: str) -> str:
    text = ITALIC_MARK.sub(lambda m: m.group(0)[1:-1], text)
    return re.sub(r"⟪(.*?)⟫", r"\1", text)
PHONETIC_FONTS = {"YuanSi_1"}
TITLE_FONTS = {"NEU-H4-Regular", "CastleT-Bold", "Georgia-Bold"}
DROP_CAP_MIN = 28
TITLE_MIN = 16.5
BODY_MIN, BODY_MAX = 9.3, 12.2
STOP_HEADINGS = {
    "Word Bank",
    "Topic Vocabulary",
    "Did You Know?",
    "Slide Show",
    "People in Focus",
    "用报说明",
    "学术指导",
    "Summary Writing",
    "Learn to Write",
    "Sentence Structures",
    "Vocabulary",
}
NOISE_LINE = re.compile(
    r"^(编辑[:：].*|高中进阶版|SHANGHAI STUDENTS.? POST|Tuesday,.*|"
    r"广告|扫码.*|主管、主办单位.*|国内统一连续出版物号.*|本期4开.*|"
    r"教材链接[:：].*|作者[:：].*|指导老师[:：].*|本版文章系.*|"
    r"A[1-8]\s+Tuesday.*|B[1-4]\s+Tuesday.*)$",
    re.I,
)
EXTRA_HEADINGS = {
    "Cloze": "完形填空",
    "完形填空": "完形填空",
    "Vocabulary": "专项词汇",
    "语法填空": "语法填空",
    "选词填空": "选词填空",
    "Grammar Filling": "语法填空",
    "Banked Cloze": "选词填空",
    "Word Filling": "语法填空",
}


@dataclass
class Span:
    x0: float
    y0: float
    x1: float
    y1: float
    size: float
    font: str
    text: str
    page: int = 0

    @property
    def mid_x(self) -> float:
        return (self.x0 + self.x1) / 2


@dataclass
class Article:
    id: str
    page: str
    title: str
    kicker: str = ""
    genre: str = ""
    words: int = 0
    theme: str = ""
    level: str = ""
    source: str = ""
    textbook: str = ""
    passage: str = ""
    passage_blocks: list[dict[str, Any]] = field(default_factory=list)
    word_bank: list[str] = field(default_factory=list)
    exercises: list[dict[str, Any]] = field(default_factory=list)
    diagrams: list[bytes] = field(default_factory=list)
    images: list[bytes] = field(default_factory=list)
    page_index: int = 0
    kind: str = "reading"
    box: tuple[float, float, float, float] | None = None


def _unique_spans(page: pymupdf.Page, page_index: int) -> list[Span]:
    data = page.get_text("dict", flags=pymupdf.TEXT_PRESERVE_WHITESPACE)
    seen: set[tuple] = set()
    spans: list[Span] = []
    for block in data.get("blocks") or []:
        if block.get("type") != 0:
            continue
        for line in block.get("lines") or []:
            for raw in line.get("spans") or []:
                text = (raw.get("text") or "").replace("\u00a0", " ").strip()
                if not text:
                    continue
                font = raw.get("font") or ""
                if font in PHONETIC_FONTS:
                    continue
                x0, y0, x1, y1 = raw["bbox"]
                key = (text, int(round(x0 / 3)), int(round(y0 / 3)), round(float(raw.get("size") or 0), 1))
                if key in seen:
                    continue
                seen.add(key)
                spans.append(
                    Span(
                        x0=x0,
                        y0=y0,
                        x1=x1,
                        y1=y1,
                        size=float(raw.get("size") or 0),
                        font=font,
                        text=text,
                        page=page_index,
                    )
                )
    return spans


def _clean_title(text: str) -> str:
    text = _ascii_quotes(re.sub(r"\s+", " ", text).strip())
    text = APOSTROPHE_RE.sub(r"\1's", text)
    words = text.split()
    n = len(words)
    for length in range(1, n // 2 + 1):
        chunk = words[:length]
        repeats = n // length
        if repeats >= 2 and words[: repeats * length] == chunk * repeats:
            return " ".join(chunk + words[repeats * length :]).strip()
    return text


def _slug(title: str) -> str:
    slug = re.sub(r"[^A-Za-z0-9]+", "-", title.lower()).strip("-")
    return slug[:40] or "article"


def _cluster(values: Iterable[float], gap: float) -> list[float]:
    ordered = sorted(values)
    if not ordered:
        return []
    groups = [[ordered[0]]]
    for value in ordered[1:]:
        if value - groups[-1][-1] > gap:
            groups.append([value])
        else:
            groups[-1].append(value)
    return [sum(g) / len(g) for g in groups]


def _lines(spans: list[Span], y_tol: float = 3.6) -> list[list[Span]]:
    if not spans:
        return []
    ordered = sorted(spans, key=lambda s: (s.y0, s.x0))
    lines: list[list[Span]] = []
    for span in ordered:
        if lines and abs(span.y0 - lines[-1][0].y0) <= y_tol:
            lines[-1].append(span)
        else:
            lines.append([span])
    for line in lines:
        line.sort(key=lambda s: s.x0)
    return lines


def _join_tokens(parts: list[str]) -> str:
    text = ""
    for part in parts:
        piece = _ascii_quotes(part.strip())
        if not piece:
            continue
        if not text:
            text = piece
            continue
        if text.endswith("-") and piece[:1].islower():
            text = text[:-1] + piece
            continue
        if piece[:1] in {"'", "’", "‘"} or text.endswith(("'", "’", "‘")):
            text += piece
            continue
        if text[-1] in "/（(" or piece[:1] in "/）),.;:!?":
            text += piece
            continue
        text += " " + piece
    text = re.sub(r"\(\s+", "(", text)
    text = re.sub(r"\s+\)", ")", text)
    text = re.sub(r"\s+([,.;:!?])", r"\1", text)
    text = APOSTROPHE_RE.sub(r"\1's", text)
    return _ascii_quotes(re.sub(r" {2,}", " ", text).strip())


def _line_text(line: list[Span], markup: bool = False) -> str:
    if not markup:
        return _join_tokens([s.text for s in line])
    ratio = _line_bold_ratio(line)
    mixed_bold = 0.05 < ratio < 0.92
    plain = _join_tokens([s.text for s in line])
    all_bold_cont = ratio >= 0.92 and bool(plain[:1].islower())
    mark_bold = mixed_bold or all_bold_cont
    parts = []
    for span in line:
        text = span.text
        font = span.font or ""
        stripped = text.strip()
        if stripped and "Italic" in font and re.search(r"[A-Za-z]{2,}", stripped):
            parts.append(f"*{stripped}*")
        elif mark_bold and stripped and "Bold" in font and re.search(r"[A-Za-z]{2,}", stripped):
            core = re.sub(r"[.,;:!?]+$", "", stripped)
            tail = stripped[len(core) :]
            parts.append(f"⟪{core}⟫{tail}" if core else text)
        else:
            parts.append(text)
    return _join_tokens(parts)


def _line_bold_ratio(line: list[Span]) -> float:
    total = sum(len(span.text.strip()) for span in line)
    if not total:
        return 0.0
    bold = sum(len(span.text.strip()) for span in line if "Bold" in (span.font or ""))
    return bold / total


def _column_lefts(spans: list[Span], gutter: float = 12) -> list[float]:
    starts: list[float] = []
    for line in _lines(spans):
        starts.append(line[0].x0)
        prev_x1 = line[0].x1
        for span in line[1:]:
            if span.x0 - prev_x1 >= gutter:
                starts.append(span.x0)
            prev_x1 = max(prev_x1, span.x1)
    lefts = _cluster(starts, gap=40)
    if len(lefts) <= 1:
        return lefts
    merged: list[float] = []
    for left in lefts:
        if merged and left - merged[-1] < 90:
            continue
        merged.append(left)
    return merged


def _assign_column(x: float, lefts: list[float]) -> int:
    if not lefts:
        return 0
    idx = 0
    for i, left in enumerate(lefts):
        if x + 8 >= left:
            idx = i
    return idx


def _assign_nearest_column(x: float, lefts: list[float]) -> int:
    if not lefts:
        return 0
    return min(range(len(lefts)), key=lambda i: abs(x - lefts[i]))


def _apply_drop_cap(text: str, drop: str) -> str:
    if not drop:
        return text
    match = re.match(r"^([①②③④⑤⑥⑦⑧⑨⑩])\s*(.*)$", text)
    circled, rest = (match.group(1), match.group(2)) if match else ("", text)
    rest = (rest or "").lstrip("*")
    if rest[:1].islower():
        rest = drop + rest
    elif rest:
        rest = drop + rest
    else:
        rest = drop
    return f"{circled} {rest}".strip() if circled else rest


def _typical_left(lines: list[list[Span]]) -> float:
    xs = sorted(line[0].x0 for line in lines if line)
    if not xs:
        return 0.0
    groups = [[xs[0]]]
    for x in xs[1:]:
        if x - groups[-1][-1] > 6:
            groups.append([x])
        else:
            groups[-1].append(x)
    candidates = sorted(groups, key=lambda g: (-len(g), g[0]))
    left = candidates[0]
    return sum(left) / len(left)


def _paragraphs_from_column(spans: list[Span], drop_cap: str = "") -> list[dict[str, str]]:
    lines = _lines(spans)
    blocks: list[dict[str, str]] = []
    current: dict[str, str] | None = None
    prev_y = None
    used_drop = False
    col_left = _typical_left(lines)

    def flush() -> None:
        nonlocal current
        if current and (current.get("text") or current.get("lead")):
            blocks.append(current)
        current = None

    for line in lines:
        raw = _line_text(line, markup=True)
        plain = _line_text(line)
        if not plain or NOISE_LINE.match(plain):
            continue
        if plain in STOP_HEADINGS or plain.startswith("▶"):
            break
        if "高中进阶版" in plain or "SPOTLIGHT" in plain or plain in {"CULTURE"}:
            continue
        if re.match(r"^语篇类型", plain):
            continue
        if re.search(r"- ?SSP$", plain) and len(plain) < 48:
            break
        if plain.strip() == "SSP":
            break
        if re.fullmatch(r"丂+", plain.replace(" ", "")):
            continue
        if _looks_like_glossary(plain):
            continue
        if len(plain.split()) == 1 and re.fullmatch(r"[A-Za-z]{3,16}", plain):
            continue
        y = line[0].y0
        gap = (y - prev_y) if prev_y is not None else 0
        is_bullet = bool(re.match(r"^[●•·]\s*", plain))
        is_heading = (
            not is_bullet
            and _line_bold_ratio(line) >= 0.7
            and 3 <= len(plain) < 72
            and re.search(r"[A-Za-z\u4e00-\u9fff]", plain)
            and not plain.endswith((",", ";", ":"))
            and not re.search(r"- ?SSP$", plain)
            and not CIRCLED_RE.search(plain[:2])
            and not _looks_like_glossary(plain)
            and not plain[:1].islower()
        )
        is_click = plain.startswith("Click:")
        numbered_head = bool(re.match(r"^\[\d+\]", plain) or re.match(r"^(A\d|N\d)\s*:", plain))
        dx = line[0].x0 - col_left
        indent_start = 10 <= dx <= 26
        rest_plain = re.sub(r"^[①②③④⑤⑥⑦⑧⑨⑩]\s*", "", plain)
        sentence_start = bool(
            CIRCLED_RE.match(plain)
            or re.match(r'^[A-Z“"‘\']', rest_plain)
            or re.match(r"^_{3,}\d", rest_plain)
        )
        if is_heading or numbered_head:
            flush()
            blocks.append({"kind": "heading", "text": plain})
            prev_y = y
            continue
        if is_bullet:
            flush()
            lead = re.sub(r"^[●•·]\s*", "", plain).strip()
            current = {"kind": "bullet", "lead": lead, "text": ""}
            prev_y = y
            continue
        if current and current.get("kind") == "bullet" and not current.get("text") and _line_bold_ratio(line) >= 0.7:
            current["lead"] = (current.get("lead") + " " + plain).strip() if current.get("lead") else plain
            prev_y = y
            continue
        new_para = False
        if current and current.get("kind") == "body":
            if is_click or (indent_start and sentence_start):
                new_para = True
        if current and current.get("kind") == "bullet":
            extra = raw
            current["text"] = _join_tokens([current["text"], extra]) if current.get("text") else extra
        elif current and current.get("kind") == "body" and not new_para:
            if current["text"] and len(current["text"]) == 1 and current["text"].isalpha() and raw[:1].islower():
                current["text"] = current["text"] + raw.lstrip("*")
            else:
                current["text"] = _join_tokens([current["text"], raw])
        else:
            flush()
            text = raw
            if drop_cap and not used_drop:
                text = _apply_drop_cap(text, drop_cap)
                used_drop = True
            current = {"kind": "indent" if is_click else "body", "text": text}
        prev_y = y
    flush()
    return [b for b in blocks if b.get("text") or b.get("lead")]


def _blocks_to_passage(blocks: list[dict[str, str]]) -> str:
    parts = []
    for block in blocks:
        if block.get("kind") == "heading":
            parts.append(_strip_markup(block.get("text") or ""))
        elif block.get("kind") == "bullet":
            lead = _strip_markup(block.get("lead") or "")
            text = _strip_markup(block.get("text") or "")
            parts.append(f"• {lead} {text}".strip())
        else:
            parts.append(_strip_markup(block.get("text") or ""))
    return "\n\n".join(p for p in parts if p)


def _is_body_span(span: Span) -> bool:
    if span.text.strip() in {"●", "•", "·"} and 7.5 <= span.size <= 14:
        return True
    return BODY_MIN <= span.size <= BODY_MAX


def _page_code(spans: list[Span], x: float) -> str:
    header = [s for s in spans if s.y0 < 58 and abs(s.mid_x - x) < 420]
    for span in header:
        if span.size >= 28 and PAGE_CODE_RE.match(span.text.strip()):
            return span.text.strip()
    blob = " ".join(s.text for s in header)
    found = re.findall(r"\b([AB][1-8])\b", blob)
    found = [c for c in found if "版" not in blob[max(0, blob.find(c) - 2) : blob.find(c) + 4]]
    return found[0] if found else ""


def _meta_blob(spans: list[Span]) -> dict[str, str]:
    text = re.sub(r"\s+", " ", " ".join(s.text for s in spans))
    match = META_RE.search(text)
    return match.groupdict() if match else {}


def _title_above(spans: list[Span], meta: Span, x1: float) -> str:
    band = [
        s
        for s in spans
        if s.size >= TITLE_MIN
        and s.y0 < meta.y0
        and s.y0 > meta.y0 - 130
        and meta.x0 - 40 <= s.x0 < x1
        and re.search(r"[A-Za-z]{2,}", s.text)
        and not (s.size >= DROP_CAP_MIN and re.fullmatch(r"[A-Za-z]", s.text))
    ]
    lines = []
    for line in _lines(band, y_tol=10):
        if abs(line[0].x0 - meta.x0) > 110:
            continue
        piece = _clean_title(_line_text(line))
        if piece and piece not in lines:
            lines.append(piece)
    return _clean_title(" ".join(lines))


def _kicker_above(spans: list[Span], meta: Span, x1: float) -> str:
    bits = []
    for span in spans:
        if 11.8 <= span.size <= 14.5 and meta.y0 - 90 < span.y0 < meta.y0 and meta.x0 - 40 <= span.x0 < x1:
            if re.search(r"[\u4e00-\u9fff]", span.text):
                bits.append(span.text.strip())
    return _join_tokens(bits)


def _drop_cap(spans: list[Span]) -> str:
    caps = [s for s in spans if s.size >= DROP_CAP_MIN and re.fullmatch(r"[A-Za-z]", s.text)]
    if not caps:
        return ""
    return sorted(caps, key=lambda s: s.y0)[0].text.upper()


def _source(spans: list[Span]) -> str:
    for span in sorted(spans, key=lambda s: s.y0, reverse=True):
        if re.search(r"-SSP$", span.text.strip()) and span.size < 16:
            return span.text.strip()
    return ""


def _word_bank(spans: list[Span]) -> list[str]:
    started = False
    items: list[str] = []
    for line in _lines([s for s in spans if s.size < 16]):
        raw = _line_text(line)
        if raw == "Word Bank":
            started = True
            continue
        if not started:
            continue
        if raw in STOP_HEADINGS or NOISE_LINE.match(raw):
            break
        cleaned = re.sub(r"/[^/]*/", "", raw).strip()
        cleaned = re.sub(r"→\s*", "→ ", cleaned)
        if re.search(r"[A-Za-z]", cleaned):
            items.append(cleaned)
    return items


GLOSS_RE = re.compile(
    r"^(?:→\s*)?(?:[A-Za-z][-A-Za-z']{0,40}\s+)?(?:n|v|adj|adv|prep|num)\.\s"
)


def _looks_like_glossary(line: str) -> bool:
    stripped = _strip_markup(line.strip())
    if GLOSS_RE.match(stripped):
        return True
    if re.match(r"^(n|v|adj|adv|prep|num)\. ", stripped) and len(stripped) < 50:
        return True
    if re.search(r"\b(n|v|adj|adv)\.\s*[\u4e00-\u9fff]", stripped) and len(stripped) < 90:
        return True
    return False


def _passage(spans: list[Span], meta_y: float) -> list[dict[str, str]]:
    body = []
    for span in spans:
        if span.y0 < meta_y + 1:
            continue
        if span.size >= TITLE_MIN:
            continue
        if not _is_body_span(span):
            continue
        if span.font in TITLE_FONTS and span.size >= 13:
            continue
        body.append(span)
    if not body:
        return []
    lefts = _column_lefts(body)
    if not lefts:
        lefts = [min(s.x0 for s in body)]
    merged: list[float] = []
    for left in lefts:
        if merged and left - merged[-1] < 220:
            continue
        merged.append(left)
    lefts = merged or lefts
    columns: list[list[Span]] = [[] for _ in lefts]
    for span in body:
        columns[_assign_column(span.x0, lefts)].append(span)
    drop = _drop_cap(spans)
    ssp_y = min(
        (
            s.y0
            for s in spans
            if s.y0 > meta_y
            and s.size < 16
            and (re.search(r"- ?SSP$", s.text.strip()) or s.text.strip() == "SSP")
        ),
        default=None,
    )
    parts: list[dict[str, str]] = []
    for i, col in enumerate(columns):
        col = _trim_column_after_source(col, ssp_y)
        blocks = [
            block
            for block in _paragraphs_from_column(col, drop if i == 0 else "")
            if not _looks_like_glossary(block.get("text") or block.get("lead") or "")
        ]
        if (
            parts
            and blocks
            and parts[-1].get("kind") == "body"
            and blocks[0].get("kind") == "body"
        ):
            lead = _strip_markup(blocks[0].get("text") or "").lstrip()
            if lead and (lead[0].islower() or lead[0] in ",.;:)"):
                parts[-1]["text"] = _join_tokens([parts[-1].get("text") or "", blocks[0].get("text") or ""])
                blocks = blocks[1:]
        parts.extend(blocks)
    return parts


def _trim_column_after_source(col: list[Span], ssp_y: float | None) -> list[Span]:
    if not col or ssp_y is None:
        return col
    above = [s for s in col if s.y0 <= ssp_y + 8]
    below = [s for s in col if s.y0 > ssp_y + 8]
    if not below or not above:
        return col
    max_above = max(s.y0 for s in above)
    min_below = min(s.y0 for s in below)
    if min_below - max_above > 32:
        return above
    return col


def _boxes_overlap_x(a0: float, a1: float, b0: float, b1: float, pad: float = 8) -> bool:
    return not (a1 <= b0 + pad or b1 <= a0 + pad)


def _article_box(
    anchors: list[Span],
    current: Span,
    page_width: float,
    page_height: float,
    page_spans: list[Span] | None = None,
) -> tuple[float, float, float, float]:
    same_page = [a for a in anchors if a.page == current.page and a is not current]
    right = min((a.x0 for a in same_page if a.x0 > current.x0 + 70), default=None)
    left_edge = max((a.x1 for a in same_page if a.x0 < current.x0 - 40), default=None)
    x0 = max(0.0, current.x0 - 24)
    if left_edge is not None:
        x0 = max(x0, left_edge + 12)
    x1 = min(right - 22, page_width) if right is not None else page_width
    overlapping = [a for a in same_page if _boxes_overlap_x(x0, x1, a.x0, a.x1)]
    below = min((a.y0 for a in overlapping if a.y0 > current.y0 + 50), default=page_height)
    if page_spans:
        next_title = [
            s.y0
            for s in page_spans
            if s.size >= TITLE_MIN
            and (s.font in TITLE_FONTS or s.size >= 22)
            and x0 - 10 <= s.x0 < x1 + 10
            and s.y0 > current.y0 + 80
            and re.search(r"[A-Za-z]{3,}", s.text)
        ]
        if next_title:
            below = min(below, min(next_title))
        prev_ssp = [
            s.y0
            for s in page_spans
            if x0 - 10 <= s.x0 < x1 + 10
            and s.y0 < current.y0 - 20
            and s.size < 16
            and (re.search(r"- ?SSP$", s.text.strip()) or s.text.strip() == "SSP")
        ]
        y0 = max(0.0, current.y0 - 140)
        if prev_ssp:
            y0 = max(y0, max(prev_ssp) + 6)
    else:
        y0 = max(0.0, current.y0 - 140)
    y1 = below - 6
    return x0, y0, x1, y1


def _in_box(span: Span, box: tuple[float, float, float, float], page: int) -> bool:
    x0, y0, x1, y1 = box
    return span.page == page and x0 - 8 <= span.mid_x <= x1 + 8 and y0 <= span.y0 <= y1


def extract_issue(path: str | Path) -> dict[str, Any]:
    doc = pymupdf.open(path)
    try:
        first_text = doc[0].get_text()
        date = ""
        m = re.search(r"(Monday|Tuesday|Wednesday|Thursday|Friday),\s+[A-Z][a-z]+\.?\s+\d+,\s+\d{4}", first_text)
        if m:
            date = m.group(0)
        issue_no = ""
        n = re.search(r"第(\d+)期", first_text)
        if n:
            issue_no = n.group(1)
        textbook = ""
        tm = re.search(r"教材链接[:：]\s*([^\n]+)", first_text + "\n" + (doc[3].get_text() if doc.page_count > 3 else ""))
        if tm:
            textbook = tm.group(1).strip()

        page_spans: list[list[Span]] = []
        page_sizes: list[tuple[float, float]] = []
        for i, page in enumerate(doc):
            page_spans.append(_unique_spans(page, i))
            page_sizes.append((page.rect.width, page.rect.height))

        anchors: list[Span] = []
        for spans in page_spans:
            for span in spans:
                if "语篇类型" in span.text and span.size < 16:
                    anchors.append(span)

        articles: list[Article] = []
        used_ids: dict[str, int] = {}
        for meta in sorted(anchors, key=lambda s: (s.page, s.y0, s.x0)):
            spans = page_spans[meta.page]
            width, height = page_sizes[meta.page]
            box = _article_box(anchors, meta, width, height, spans)
            local = [s for s in spans if _in_box(s, box, meta.page)]
            title = _title_above(spans, meta, box[2])
            if not title:
                continue
            info = _meta_blob(local)
            art_id = _slug(title)
            used_ids[art_id] = used_ids.get(art_id, 0) + 1
            if used_ids[art_id] > 1:
                art_id = f"{art_id}-{used_ids[art_id]}"
            page_code = _page_code(spans, meta.x0) or _page_code(local, meta.x0)
            passage_blocks = _passage(local, meta.y1)
            articles.append(
                Article(
                    id=art_id,
                    page=page_code,
                    title=title,
                    kicker=_kicker_above(spans, meta, box[2]),
                    genre=info.get("genre", ""),
                    words=int(info.get("words") or 0),
                    theme=info.get("theme", ""),
                    level=info.get("level", ""),
                    source=_source(local),
                    textbook=textbook,
                    passage=_blocks_to_passage(passage_blocks),
                    passage_blocks=passage_blocks,
                    word_bank=_word_bank(local),
                    page_index=meta.page,
                    kind="reading",
                    box=box,
                )
            )

        debates = _extract_debates(page_spans)
        existing = {a.title.lower() for a in articles}
        for debate in debates:
            if debate.title.lower() not in existing:
                articles.append(debate)

        extras = _extract_extras(page_spans, page_sizes)
        existing_ids = {a.id for a in articles}
        articles.extend([a for a in extras if a.id not in existing_ids])

        _attach_exercises(articles, page_spans, page_sizes)
        _attach_diagrams(articles, doc)
        return {
            "paper": "Shanghai Students' Post 高中进阶版",
            "date": date,
            "issue": issue_no,
            "source_pages": doc.page_count,
            "articles": [asdict(a) for a in articles],
        }
    finally:
        doc.close()


def _extract_debates(page_spans: list[list[Span]]) -> list[Article]:
    found: list[Article] = []
    for spans in page_spans:
        titles = [
            s
            for s in spans
            if s.size >= 22 and (s.font in TITLE_FONTS or s.size >= 24)
            and re.search(r"[A-Za-z]{4,}", s.text)
        ]
        if not titles:
            continue
        grouped: list[list[Span]] = []
        for span in sorted(titles, key=lambda s: (s.y0, s.x0)):
            if grouped and abs(span.y0 - grouped[-1][-1].y0) < 40 and abs(span.x0 - grouped[-1][0].x0) < 30:
                grouped[-1].append(span)
            else:
                grouped.append([span])
        for group in grouped:
            title = _clean_title(" ".join(_line_text(line) for line in _lines(group)))
            if not title.lower().startswith("should"):
                continue
            y0 = min(s.y0 for s in group)
            y1 = 1010
            local = [s for s in spans if s.y0 > y0 - 10 and s.y0 < y1 and s.x0 >= min(t.x0 for t in group) - 20]
            # cut at next debate title
            others = [g for g in grouped if min(s.y0 for s in g) > y0 + 30]
            if others:
                y1 = min(s.y0 for s in others[0]) - 8
                local = [s for s in local if s.y0 < y1]
            body = [s for s in local if BODY_MIN <= s.size <= BODY_MAX]
            lefts = _column_lefts(body) or [min((s.x0 for s in body), default=0)]
            cols: list[list[Span]] = [[] for _ in lefts]
            for span in body:
                cols[_assign_nearest_column(span.x0, lefts)].append(span)
            parts: list[dict[str, str]] = []
            for col in cols:
                parts.extend(_paragraphs_from_column(col))
            passage = _blocks_to_passage(parts).strip()
            if len(passage) < 120:
                continue
            found.append(
                Article(
                    id=_slug(title),
                    page=_page_code(spans, group[0].x0),
                    title=title,
                    kicker="",
                    genre="辩论",
                    passage=passage,
                    kind="debate",
                )
            )
    return found


def _extra_heading_spans(spans: list[Span]) -> list[Span]:
    found: list[Span] = []
    for span in spans:
        text = span.text.strip()
        if text not in EXTRA_HEADINGS:
            continue
        if span.size < 12.5:
            continue
        found.append(span)
    found.sort(key=lambda s: (s.y0, s.x0))
    return found


def _extra_subtitle(local: list[Span], head: Span, *, after_y: float | None = None) -> tuple[str, float]:
    y_min = head.y1 if after_y is None else after_y
    band = [
        s
        for s in local
        if "Bold" in (s.font or "")
        and y_min <= s.y0 <= y_min + 50
        and 8 <= len(s.text.strip()) <= 60
        and re.search(r"[A-Za-z]{3,}", s.text)
        and not re.search(r"_{2,}|\(\d+\)", s.text)
        and not NUMBERED_RE.match(s.text.strip())
        and not LABELED_WORD_RE.match(s.text.strip())
    ]
    if not band:
        return "", y_min
    y0 = min(s.y0 for s in band)
    line = [s for s in band if abs(s.y0 - y0) < 6]
    title = _clean_title(_join_tokens(s.text for s in sorted(line, key=lambda s: s.x0)))
    if title.endswith((".", ",", ";", ":")):
        return "", y_min
    return title, max(s.y1 for s in line)


def _labeled_word_bank(spans: list[Span], head: Span) -> tuple[list[str], float]:
    by_label: dict[str, tuple[float, str]] = {}
    y1 = head.y1
    for span in spans:
        match = LABELED_WORD_RE.match(span.text.strip())
        if not match:
            continue
        if span.y0 < head.y1 - 4 or span.y0 > head.y0 + 80:
            continue
        label, word = match.group(1), match.group(2)
        prev = by_label.get(label)
        if prev is None or span.x0 < prev[0]:
            by_label[label] = (span.x0, word)
        y1 = max(y1, span.y1)
    if len(by_label) < 8:
        return [], head.y1
    items = [f"{label}. {by_label[label][1]}" for label in "ABCDEFGHIJK" if label in by_label]
    return items, y1


def _extract_extras(
    page_spans: list[list[Span]],
    page_sizes: list[tuple[float, float]],
) -> list[Article]:
    found: list[Article] = []
    seen: set[str] = set()
    for page_index, (spans, (width, height)) in enumerate(zip(page_spans, page_sizes)):
        heads = _extra_heading_spans(spans)
        if not heads:
            continue
        mid = width / 2
        for i, head in enumerate(heads):
            hx0 = 0.0 if head.mid_x < mid else mid
            hx1 = mid if head.mid_x < mid else width
            y1 = height - 4
            for other in heads[i + 1 :]:
                if other.y0 > head.y0 + 24 and not (other.x0 >= hx1 or other.x1 <= hx0):
                    y1 = min(y1, other.y0 - 6)
            local = [
                s
                for s in spans
                if hx0 - 4 <= s.mid_x <= hx1 + 4 and head.y1 < s.y0 <= y1
            ]
            if not local:
                continue
            bank, bank_y1 = _labeled_word_bank(local, head)
            passage_spans = [
                s
                for s in local
                if not (bank and LABELED_WORD_RE.match(s.text.strip()) and s.y0 <= bank_y1 + 4)
            ]
            subtitle, sub_y1 = _extra_subtitle(passage_spans, head, after_y=bank_y1 if bank else None)
            lines = _merge_circled_lines(
                _collect_box_lines(passage_spans, hx0, sub_y1 + 1, hx1, y1, min_col_gap=220)
            )
            if subtitle and lines:
                if lines[0].strip() == subtitle or subtitle in lines[0]:
                    lines = lines[1:]
            if len(" ".join(lines)) < 80:
                continue
            quiz_at = _quiz_line_index(lines)
            title_lines = lines if quiz_at is None else lines[:quiz_at]
            quiz_lines = [] if quiz_at is None else lines[quiz_at:]
            paras = _lines_to_paragraphs(title_lines)
            passage_blocks = [{"kind": "body", "text": p} for p in paras if p]
            exercises = _parse_exercises(quiz_lines)
            genre = EXTRA_HEADINGS.get(head.text.strip(), "专项")
            if bank:
                genre = "选词填空"
            title = subtitle or head.text.strip()
            art_id = _slug(f"{head.text.strip()}-{title}")
            if art_id in seen:
                continue
            seen.add(art_id)
            found.append(
                Article(
                    id=art_id,
                    page=_page_code(spans, head.x0) or ("B3" if head.mid_x >= mid else "B2"),
                    title=title,
                    kicker=head.text.strip(),
                    genre=genre,
                    passage=_blocks_to_passage(passage_blocks),
                    passage_blocks=passage_blocks,
                    word_bank=bank,
                    exercises=exercises,
                    page_index=page_index,
                    kind="extra",
                )
            )
    return found


def _quiz_line_index(lines: list[str]) -> int | None:
    for i, line in enumerate(lines):
        stripped = line.strip()
        if re.match(r"^\d+[\.、．]\s*A[\.、．]", stripped):
            return i
        if SECTION_RE.match(stripped) or stripped.startswith("I.") or stripped.startswith("II.") or stripped.startswith("III."):
            return i
    return None


def _lines_to_paragraphs(lines: list[str]) -> list[str]:
    paras: list[str] = []
    buf: list[str] = []
    for line in lines:
        piece = line.strip()
        if not piece:
            continue
        buf.append(piece)
        if piece.endswith((".", "?", "!", "。")) and len(piece) < 78:
            paras.append(_join_tokens(buf))
            buf = []
    if buf:
        paras.append(_join_tokens(buf))
    return paras


def _merge_circled_lines(lines: list[str]) -> list[str]:
    out: list[str] = []
    i = 0
    while i < len(lines):
        cur = lines[i].strip()
        if CIRCLED_ONLY_RE.match(cur) and i + 1 < len(lines):
            nxt = lines[i + 1].strip()
            if nxt and not CIRCLED_ONLY_RE.match(nxt):
                out.append(f"{cur} {nxt}")
                i += 2
                continue
        out.append(lines[i])
        i += 1
    return out


def _quiz_headers(spans: list[Span]) -> list[tuple[Span, str]]:
    codes = [
        s
        for s in spans
        if re.fullmatch(r"A[1-8]", s.text.strip())
        and 9.0 <= s.size <= 14.5
        and s.y0 > 70
    ]
    found: list[tuple[Span, str]] = []
    for code in sorted(codes, key=lambda s: (s.y0, s.x0)):
        next_code_x = min(
            (
                s.x0
                for s in codes
                if s is not code and abs(s.y0 - code.y0) <= 5 and s.x0 > code.x0
            ),
            default=code.x0 + 430,
        )
        title_bits = [
            s
            for s in spans
            if abs(s.y0 - code.y0) <= 5
            and s.x0 > code.x1 - 4
            and s.x0 < next_code_x - 4
            and 9.0 <= s.size <= 14.5
            and not re.fullmatch(r"A[1-8]", s.text.strip())
        ]
        title = _clean_title(_join_tokens([s.text for s in sorted(title_bits, key=lambda s: s.x0)]))
        if len(title) >= 8 and re.search(r"[A-Za-z]{4,}", title):
            found.append((code, title))
    return found


def _is_teaching_notes(spans: list[Span]) -> bool:
    return any(
        s.text.strip() in {"Sentence Structures", "Vocabulary"} and s.size >= 11.5 and s.y0 < 220
        for s in spans
    )


def _stop_exercise_line(raw: str) -> bool:
    stripped = raw.strip()
    if stripped in STOP_HEADINGS or stripped.startswith("▶"):
        return True
    return stripped.startswith("Learn to Write") or stripped.startswith("Letter From")


def _collect_box_lines(
    spans: list[Span],
    x0: float,
    y0: float,
    x1: float,
    y1: float,
    *,
    min_col_gap: float = 140,
) -> list[str]:
    local = [
        s
        for s in spans
        if x0 - 6 <= s.mid_x <= x1 + 6 and y0 - 2 <= s.y0 <= y1 and 9.2 <= s.size <= 13.5
    ]
    if not local:
        return []
    if (x1 - x0) <= 420:
        lefts = [min(s.x0 for s in local)]
    else:
        lefts = _column_lefts(local, gutter=16) or [min(s.x0 for s in local)]
        merged: list[float] = []
        for left in lefts:
            if merged and left - merged[-1] < min_col_gap:
                continue
            merged.append(left)
        lefts = merged
    cols: list[list[Span]] = [[] for _ in lefts]
    for span in local:
        cols[_assign_nearest_column(span.x0, lefts)].append(span)

    if _looks_like_mindmap(local):
        return _mindmap_lines(local, cols)

    lines: list[str] = []
    for col in cols:
        for line in _lines(col, y_tol=5.2):
            raw = _line_text(line, markup=True)
            if not raw or NOISE_LINE.match(_strip_markup(raw)) or _stop_exercise_line(_strip_markup(raw)):
                if _strip_markup(raw) in STOP_HEADINGS or (raw and raw.startswith("Learn to Write")):
                    return lines
                continue
            if re.fullmatch(r"丂+", raw.replace(" ", "")):
                continue
            if ARTICLE_HEADER_RE.match(_strip_markup(raw).strip()):
                continue
            lines.append(raw)
    return _merge_circled_lines(lines)


def _looks_like_mindmap(spans: list[Span]) -> bool:
    blob = " ".join(s.text for s in spans)
    return "mind map" in blob.lower() and bool(re.search(r"\([1-6]\)_{3,}", blob))


def _mindmap_lines(spans: list[Span], cols: list[list[Span]]) -> list[str]:
    lines: list[str] = []
    intro = [
        _line_text(line)
        for line in _lines(spans, y_tol=3.2)
        if _line_text(line).startswith("I.") or _line_text(line).startswith("II.")
    ]
    if intro:
        lines.append(intro[0])
    boxes: list[str] = []
    for col in cols:
        parts = []
        for line in _lines(col, y_tol=3.4):
            raw = _line_text(line)
            if not raw or "丂" in raw or NOISE_LINE.match(raw) or _stop_exercise_line(raw):
                continue
            if raw.startswith("I.") or raw.startswith("II."):
                continue
            if ARTICLE_HEADER_RE.match(raw.strip()):
                continue
            parts.append(raw)
        blob = re.sub(r"_{8,}", "", " ".join(parts))
        blob = re.sub(r"\s+", " ", blob).strip()
        if re.search(r"\(\d+\)_{3,}", blob):
            boxes.append(blob)
    seen: set[str] = set()
    for box in boxes:
        key = re.search(r"\((\d+)\)", box)
        mark = key.group(1) if key else box[:40]
        if mark in seen:
            continue
        seen.add(mark)
        lines.append(box)
    for title in intro[1:]:
        lines.append(title)
        lines.extend(["_" * 72, "_" * 72, "_" * 72])
        break
    return lines


def _attach_exercises(
    articles: list[Article],
    page_spans: list[list[Span]],
    page_sizes: list[tuple[float, float]],
) -> None:
    for spans, (width, height) in zip(page_spans, page_sizes):
        mid = width / 2
        halves = (
            [s for s in spans if s.mid_x < mid - 8],
            [s for s in spans if s.mid_x >= mid - 8],
        )
        bounds = ((0.0, mid), (mid, width))
        for local, (hx0, hx1) in zip(halves, bounds):
            if not local or _is_teaching_notes(local):
                continue
            headers = _quiz_headers(local)
            if not headers:
                continue
            for i, (code, title) in enumerate(headers):
                article = _match_article(title, articles)
                if article is None:
                    continue
                next_right = [
                    h[0].x0
                    for h in headers
                    if h[0] is not code and h[0].x0 > code.x0 + 40 and abs(h[0].y0 - code.y0) < 30
                ]
                next_below = [
                    h[0].y0
                    for h in headers
                    if h[0].y0 > code.y0 + 40 and not (h[0].x0 >= (min(next_right) if next_right else hx1) or h[0].x1 <= code.x0)
                ]
                x0 = max(hx0, code.x0 - 12)
                x1 = min(next_right) - 8 if next_right else hx1
                y1 = min(next_below) - 6 if next_below else height - 8
                for span in local:
                    if span.text.strip() in STOP_HEADINGS and span.y0 > code.y0 + 20 and span.y0 < y1:
                        if x0 - 20 <= span.x0 <= x1 + 20:
                            y1 = min(y1, span.y0 - 4)
                lines = _collect_box_lines(local, x0, code.y0 + 1, x1, y1)
                if not lines:
                    continue
                article.exercises.extend(_trim_title_leftovers(_parse_exercises(lines), article))


def _attach_diagrams(articles: list[Article], doc: pymupdf.Document) -> None:
    for article in articles:
        if not _needs_diagram(article.exercises):
            continue
        if article.page_index < 0 or article.page_index >= doc.page_count:
            continue
        page = doc[article.page_index]
        infos = page.get_image_info() or []
        boxes = []
        for info in infos:
            bbox = info.get("bbox")
            if not bbox:
                continue
            rect = pymupdf.Rect(bbox)
            if rect.width >= 180 and rect.height >= 140:
                boxes.append(rect)
        if not boxes:
            continue
        unique: list[pymupdf.Rect] = []
        for rect in sorted(boxes, key=lambda r: r.y0):
            if any(abs(rect.x0 - u.x0) < 12 and abs(rect.y0 - u.y0) < 12 for u in unique):
                continue
            unique.append(rect)
        best = max(unique, key=lambda r: r.width * r.height)
        pix = page.get_pixmap(clip=best, matrix=pymupdf.Matrix(2.0, 2.0), alpha=False)
        article.diagrams.append(pix.tobytes("png"))


def _needs_diagram(exercises: list[dict[str, Any]]) -> bool:
    for block in exercises:
        for question in block.get("questions") or []:
            if re.search(r"\bdiagram\b", str(question.get("stem") or ""), re.I):
                return True
        if re.search(r"\bdiagram\b", str(block.get("title") or ""), re.I):
            return True
    return False


def _norm_title(text: str) -> str:
    text = text.replace("’", "'").replace("`", "'")
    return re.sub(r"[^a-z0-9]+", " ", text.lower()).strip()


def _match_article(fragment: str, articles: list[Article]) -> Article | None:
    needle = _norm_title(fragment)
    if not needle:
        return None
    scored: list[tuple[int, Article]] = []
    for article in articles:
        title = _norm_title(article.title)
        if not title:
            continue
        if title == needle or needle.startswith(title) or title.startswith(needle):
            scored.append((100 + min(len(title), len(needle)), article))
            continue
        tw, nw = title.split(), needle.split()
        common = 0
        for left, right in zip(tw, nw):
            if left == right:
                common += 1
            else:
                break
        if common >= 3:
            scored.append((common * 10, article))
    if not scored:
        return None
    scored.sort(key=lambda item: -item[0])
    return scored[0][1]


def _trim_title_leftovers(blocks: list[dict[str, Any]], article: Article) -> list[dict[str, Any]]:
    needle = _norm_title(article.title)
    kept: list[dict[str, Any]] = []
    for block in blocks:
        lines = []
        for line in block.get("lines") or []:
            norm = _norm_title(_strip_markup(str(line)))
            if norm and needle and norm in needle:
                continue
            lines.append(line)
        block["lines"] = lines
        if block.get("title") or block.get("questions") or lines:
            kept.append(block)
    return kept


def _parse_exercises(lines: list[str]) -> list[dict[str, Any]]:
    blocks: list[dict[str, Any]] = []
    current: dict[str, Any] | None = None
    question: dict[str, Any] | None = None

    def flush_question() -> None:
        nonlocal question
        if not question or current is None:
            question = None
            return
        if question.get("options"):
            current.setdefault("questions", []).append(question)
        else:
            stem = f"{question.get('no')}. {question.get('stem') or ''}".strip()
            current.setdefault("lines", []).append(stem)
        question = None

    def flush_block() -> None:
        nonlocal current
        flush_question()
        if current:
            blocks.append(current)
            current = None

    for line in _merge_circled_lines(lines):
        if SECTION_RE.match(line) or line.startswith("I.") or line.startswith("II.") or line.startswith("III."):
            flush_block()
            current = {"kind": "section", "title": line, "questions": [], "lines": []}
            continue
        if current is None:
            current = {"kind": "section", "title": "", "questions": [], "lines": []}
        if (
            current["title"]
            and not current["questions"]
            and not current["lines"]
            and question is None
            and _is_title_wrap(line)
        ):
            current["title"] = f"{current['title']} {line}".strip()
            continue
        numbered = NUMBERED_RE.match(line)
        choice_parts = _choice_parts(line)
        if numbered:
            flush_question()
            stem = numbered.group(2).strip()
            question = {"no": numbered.group(1), "stem": stem, "options": []}
            inline = _choice_parts(stem)
            if inline:
                question["stem"] = ""
                question["options"].extend(inline)
            continue
        if choice_parts and question is not None:
            question["options"].extend(choice_parts)
            continue
        if question is None and not _is_circled_item(line) and not _is_word_bank_line(line):
            last_opt = bool(current.get("lines") and _is_choice_line(str(current["lines"][-1])))
            if choice_parts or last_opt:
                for chunk in _split_choice_chunks(line):
                    if _is_choice_line(chunk):
                        current.setdefault("lines", []).append(chunk)
                    elif current.get("lines") and _is_choice_line(str(current["lines"][-1])):
                        current["lines"][-1] = _join_tokens([str(current["lines"][-1]), chunk])
                    else:
                        current.setdefault("lines", []).append(chunk)
                continue
        if question is not None and not choice_parts:
            if question["options"]:
                question["options"][-1]["text"] += " " + line
            else:
                question["stem"] += " " + line
            continue
        current.setdefault("lines", []).append(line)
    flush_block()
    return [b for b in blocks if b.get("title") or b.get("questions") or b.get("lines")]


def _split_choice_chunks(line: str) -> list[str]:
    return [part.strip() for part in re.split(r"\s+(?=[A-F][\.、．]\s)", line.strip()) if part.strip()]


def _choice_parts(line: str) -> list[dict[str, str]]:
    stripped = line.strip()
    if not CHOICE_RE.match(stripped):
        return []
    chunks = re.split(r"\s+(?=[A-F][\.、．]\s)", stripped)
    parts: list[dict[str, str]] = []
    for chunk in chunks:
        match = CHOICE_RE.match(chunk.strip())
        if match:
            parts.append({"label": match.group(1), "text": match.group(2).strip()})
    return parts


def _is_title_wrap(line: str) -> bool:
    stripped = line.strip()
    if not stripped or stripped.startswith("_") or NUMBERED_RE.match(stripped) or CHOICE_RE.match(stripped):
        return False
    if re.match(r"^[①②③④⑤⑥⑦⑧⑨⑩]", stripped):
        return False
    words = stripped.split()
    if 1 <= len(words) <= 6 and all(re.fullmatch(r"[A-Za-z’-]+", w) for w in words):
        return False
    return stripped[:1].islower() or stripped.startswith(("proper", "then complete", "from the box", "in their"))


READING_HINT = re.compile(
    r"choose the best answer|mind map|summary|blank|guess the meaning|fill in each blank",
    re.I,
)
VOCAB_HINT = re.compile(
    r"translate|proper forms|given words|different forms|expressions into",
    re.I,
)


def describe_exercises(article: dict[str, Any] | Article) -> str:
    data = asdict(article) if isinstance(article, Article) else article
    bits: list[str] = []
    mcq = 0
    has_gap = False
    has_map = False
    has_vocab = False
    for block in data.get("exercises") or []:
        title = str(block.get("title") or "")
        mcq += len([q for q in block.get("questions") or [] if q.get("options")])
        if re.search(r"blank|fill in", title, re.I) or any(
            re.match(r"^[A-F]\.", ln) for ln in block.get("lines") or []
        ):
            if "blank" in title.lower() or "fill in" in title.lower():
                has_gap = True
        if "mind map" in title.lower() or "summary" in title.lower():
            has_map = True
        if VOCAB_HINT.search(title):
            has_vocab = True
    if data.get("kind") == "extra" and data.get("genre"):
        bits.insert(0, str(data["genre"]))
    if _is_labeled_word_bank(data.get("word_bank") or []):
        if "选词填空" not in bits:
            bits.append("选词填空")
    if mcq:
        bits.append(f"{mcq}道选择")
    if has_gap:
        bits.append("六选四")
    if has_map:
        bits.append("导图/摘要")
    if has_vocab:
        bits.append("词汇题")
    return "、".join(bits) or "无配套题"


def _is_labeled_word_bank(items: list[str]) -> bool:
    if len(items) < 8:
        return False
    return all(re.match(r"^[A-K]\.\s+\S+", str(item)) for item in items)


def _is_vocab_block(block: dict[str, Any]) -> bool:
    title = str(block.get("title") or "")
    if READING_HINT.search(title):
        return False
    return bool(VOCAB_HINT.search(title))


def _is_reading_block(block: dict[str, Any]) -> bool:
    title = str(block.get("title") or "")
    if READING_HINT.search(title) or block.get("questions"):
        return True
    lines = block.get("lines") or []
    return any(re.match(r"^[A-F]\. ", ln) for ln in lines) and "blank" in title.lower()


CLOZE_BLANK_RE = re.compile(r"\((\d+)\)(_{2,})")
SIX_BLANK_RE = re.compile(r"_{2,}(\d+)_{2,}")


def _af_options(block: dict[str, Any]) -> list[dict[str, str]]:
    by_label: dict[str, dict[str, str]] = {}
    for line in block.get("lines") or []:
        for part in _choice_parts(_strip_markup(str(line))):
            by_label[part["label"]] = part
    return [by_label[label] for label in "ABCDEF" if label in by_label]


def _is_six_choose_block(block: dict[str, Any]) -> bool:
    options = _af_options(block)
    if len(options) < 5:
        return False
    title = str(block.get("title") or "").lower()
    if any(q.get("options") for q in block.get("questions") or []):
        return False
    return "blank" in title or "sentence" in title or "from the box" in title


def _six_blank_ids(article: dict[str, Any]) -> list[str]:
    texts = [str(article.get("passage") or "")]
    for block in article.get("passage_blocks") or []:
        texts.append(str(block.get("text") or ""))
        texts.append(str(block.get("lead") or ""))
    blob = "\n".join(texts)
    nums = sorted({int(num) for num in SIX_BLANK_RE.findall(blob)})
    if nums:
        return [str(num) for num in nums]
    return ["1", "2", "3", "4"]


def _mcq_record(
    *,
    orig_no: str,
    article: dict[str, Any],
    block: dict[str, Any],
    stem: str,
    options: list[dict[str, str]],
    kind: str,
) -> dict[str, Any]:
    return {
        "no": 0,
        "orig_no": orig_no,
        "article_id": article.get("id") or "",
        "article_title": article.get("title") or "",
        "block_title": block.get("title") or "",
        "section": _section_code(block.get("title") or ""),
        "stem": stem,
        "options": options,
        "is_cloze": kind == "cloze",
        "kind": kind,
        "answer": "",
    }


def _remap_cloze_blanks(text: str, mapping: dict[str, int]) -> str:
    if not text or not mapping:
        return text

    def repl(match: re.Match[str]) -> str:
        new = mapping.get(match.group(1))
        if new is None:
            return match.group(0)
        return f"({new}){match.group(2)}"

    return CLOZE_BLANK_RE.sub(repl, text)


def _section_code(title: str) -> str:
    match = SECTION_RE.match((title or "").strip())
    return match.group(1) if match else ""


def _should_export_block(
    article: dict[str, Any],
    block: dict[str, Any],
    *,
    include_vocab_exercises: bool,
    include_questions: bool,
) -> bool:
    is_extra = article.get("kind") == "extra"
    if not include_questions and not is_extra:
        return False
    vocab = _is_vocab_block(block)
    if vocab and not include_vocab_exercises and not is_extra:
        return False
    if not vocab and not _is_reading_block(block) and not include_vocab_exercises and not is_extra:
        return False
    return True


def collect_mcqs(
    issue: dict[str, Any],
    article_ids: list[str],
    *,
    include_vocab_exercises: bool = True,
    include_questions: bool = True,
) -> list[dict[str, Any]]:
    lookup = {a["id"]: a for a in issue.get("articles") or []}
    ordered = [lookup[i] for i in article_ids if i in lookup]
    _, mcqs = _prepare_export_articles(
        ordered,
        include_vocab_exercises=include_vocab_exercises,
        include_questions=include_questions,
    )
    return mcqs


def _prepare_export_articles(
    ordered: list[dict[str, Any]],
    *,
    include_vocab_exercises: bool,
    include_questions: bool,
) -> tuple[list[dict[str, Any]], list[dict[str, Any]]]:
    prepared: list[dict[str, Any]] = []
    regular: list[tuple[dict[str, Any], dict[str, Any]]] = []
    six_items: list[dict[str, Any]] = []
    cloze_items: list[tuple[dict[str, Any], dict[str, Any], dict[str, Any], str]] = []
    for article in ordered:
        art = dict(article)
        art["passage_blocks"] = [dict(block) for block in article.get("passage_blocks") or []]
        art["exercises"] = []
        is_extra = article.get("kind") == "extra"
        if not include_questions and not is_extra:
            prepared.append(art)
            continue
        for block in article.get("exercises") or []:
            if not _should_export_block(
                article,
                block,
                include_vocab_exercises=include_vocab_exercises,
                include_questions=include_questions,
            ):
                continue
            copied = dict(block)
            if _is_six_choose_block(copied):
                options = _af_options(copied)
                for orig in _six_blank_ids(art):
                    six_items.append(
                        _mcq_record(
                            orig_no=orig,
                            article=art,
                            block=copied,
                            stem=f"___ {orig} ___",
                            options=options,
                            kind="six",
                        )
                    )
                art["exercises"].append(copied)
                continue
            option_qs = [q for q in (block.get("questions") or []) if q.get("options")]
            is_cloze = bool(option_qs) and _is_cloze_choice_set(option_qs)
            new_qs: list[dict[str, Any]] = []
            for question in block.get("questions") or []:
                question = dict(question)
                if question.get("options"):
                    old = str(question.get("no") or "")
                    rec = _mcq_record(
                        orig_no=old,
                        article=art,
                        block=copied,
                        stem=str(question.get("stem") or ""),
                        options=list(question.get("options") or []),
                        kind="cloze" if is_cloze else "mcq",
                    )
                    if is_cloze:
                        cloze_items.append((rec, question, art, old))
                    else:
                        regular.append((rec, question))
                new_qs.append(question)
            copied["questions"] = new_qs
            art["exercises"].append(copied)
        prepared.append(art)

    mcqs: list[dict[str, Any]] = []
    mcq_no = 1
    cloze_maps: dict[str, dict[str, int]] = {}
    for rec, question in regular:
        rec["no"] = mcq_no
        question["no"] = mcq_no
        mcqs.append(rec)
        mcq_no += 1
    for rec in six_items:
        rec["no"] = mcq_no
        mcqs.append(rec)
        mcq_no += 1
    for rec, question, art, old in cloze_items:
        rec["no"] = mcq_no
        question["no"] = mcq_no
        if old:
            cloze_maps.setdefault(str(art.get("id") or ""), {})[old] = mcq_no
        mcqs.append(rec)
        mcq_no += 1
    for art in prepared:
        mapping = cloze_maps.get(str(art.get("id") or ""))
        if not mapping:
            continue
        art["passage"] = _remap_cloze_blanks(str(art.get("passage") or ""), mapping)
        for block in art["passage_blocks"]:
            if "text" in block:
                block["text"] = _remap_cloze_blanks(str(block.get("text") or ""), mapping)
            if "lead" in block:
                block["lead"] = _remap_cloze_blanks(str(block.get("lead") or ""), mapping)
    return prepared, mcqs


def export_reading_docx(
    issue: dict[str, Any],
    article_ids: list[str],
    output_path: str | Path,
    *,
    include_word_bank: bool = False,
    include_vocab_exercises: bool = True,
    include_questions: bool = True,
) -> Path:
    lookup = {a["id"]: a for a in issue.get("articles") or []}
    ordered = [lookup[i] for i in article_ids if i in lookup]
    if not ordered:
        raise ValueError("没有选中的文章。")
    ordered, _ = _prepare_export_articles(
        ordered,
        include_vocab_exercises=include_vocab_exercises,
        include_questions=include_questions,
    )

    doc = Document()
    for section in doc.sections:
        section.page_width = Cm(21.0)
        section.page_height = Cm(29.7)
        section.top_margin = Cm(2.54)
        section.bottom_margin = Cm(2.54)
        section.left_margin = Cm(2.54)
        section.right_margin = Cm(2.54)

    style = {
        "en": "Times New Roman",
        "zh": "宋体",
        "size": 10.5,
        "line": 1.2,
        "indent": 0.74,
        "after": 0,
    }
    date = issue.get("date") or ""
    issue_no = issue.get("issue") or ""
    title = "高二英语阅读"
    sub_bits = ["Shanghai Students' Post 高中进阶版"]
    if issue_no:
        sub_bits.append(f"第{issue_no}期")
    if date:
        sub_bits.append(date)

    _para(doc, title, style, bold=True, center=True, size=16, after=6)
    _para(doc, "  ".join(sub_bits), style, italic=True, center=True, after=4)
    _para(doc, "Class: ________    Name: ________", style, center=True, after=10)

    for idx, article in enumerate(ordered):
        _para(doc, f"Passage {idx + 1}", style, bold=True, center=True, after=2)
        _para(doc, article["title"], style, bold=True, center=True, after=6)
        if article.get("kind") == "extra" and article.get("kicker") and article["kicker"] != article["title"]:
            kicker = str(article["kicker"])
            if kicker not in {"Vocabulary", "选词填空", "Banked Cloze"}:
                _para(doc, kicker, style, italic=True, center=True, after=4)

        if _is_labeled_word_bank(article.get("word_bank") or []):
            _render_vocab_option_table(doc, article["word_bank"], style)

        blocks = list(article.get("passage_blocks") or [])
        if blocks:
            _render_passage_blocks(doc, blocks, style)
        else:
            for para in str(article.get("passage") or "").split("\n\n"):
                text = para.strip()
                if not text:
                    continue
                heading = bool(re.match(r"^\[\d+\]", text) or re.match(r"^(A\d|N\d)\s*:", text) or re.match(r"^_{3,}", text))
                _para(doc, _strip_markup(text), style, justify=True, indent=not heading, after=0)

        spacer = doc.add_paragraph()
        set_paragraph_format(spacer, line_spacing=style["line"], after_pt=8)

        if include_word_bank and article.get("word_bank") and not _is_labeled_word_bank(article.get("word_bank") or []):
            _para(doc, "Word Bank", style, bold=True, after=4)
            for item in article["word_bank"]:
                _para(doc, item, style, after=0)

        is_extra = article.get("kind") == "extra"
        if not include_questions and not is_extra:
            continue

        diagrams = list(article.get("diagrams") or [])
        diagram_used = False
        for block in article.get("exercises") or []:
            vocab = _is_vocab_block(block)
            if vocab and not include_vocab_exercises and not is_extra:
                continue
            if not vocab and not _is_reading_block(block) and not include_vocab_exercises and not is_extra:
                continue
            if block.get("title"):
                title_para = _para(doc, block["title"], style, bold=True, after=4)
                title_para.paragraph_format.keep_with_next = True
            questions = [q for q in block.get("questions") or [] if q.get("options")]
            if questions and _is_cloze_choice_set(questions):
                _render_cloze_choice_table(doc, questions, style)
            elif questions:
                for question in questions:
                    stem = f"{question.get('no')}. {question.get('stem') or ''}".strip()
                    _para(doc, stem, style, justify=True, after=2)
                    if (
                        not diagram_used
                        and diagrams
                        and re.search(r"\bdiagram\b", stem, re.I)
                    ):
                        _add_diagram(doc, diagrams[0], style)
                        diagram_used = True
                    _render_choices(doc, question.get("options") or [], style)
            _render_exercise_lines(doc, block.get("lines") or [], style)
            extra = doc.add_paragraph()
            set_paragraph_format(extra, line_spacing=style["line"], after_pt=6)

    out = Path(output_path)
    out.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(out))
    return out


def _is_circled_item(line: str) -> bool:
    return bool(re.match(r"^[①②③④⑤⑥⑦⑧⑨⑩]", line.strip()))


def _write_runs(
    para,
    text: str,
    style: dict[str, Any],
    *,
    bold: bool = False,
    italic: bool = False,
    size: float | None = None,
) -> None:
    if not text:
        return
    text = _ascii_quotes(text)
    parts = re.split(r"(\*[^*]+\*|⟪[^⟫]+⟫)", text)
    for part in parts:
        if not part:
            continue
        italic_mark = part.startswith("*") and part.endswith("*") and len(part) > 2
        under_mark = part.startswith("⟪") and part.endswith("⟫") and len(part) > 2
        inner = part[1:-1] if italic_mark or under_mark else part
        add_mixed_runs(
            para,
            inner,
            en=style["en"],
            zh=style["zh"],
            size=size or style["size"],
            bold=True if under_mark else bold,
            italic=True if italic_mark else italic,
            underline=under_mark,
        )


def _add_diagram(doc: Document, png: bytes, style: dict[str, Any], *, width_cm: float = 14.2) -> None:
    pic = doc.add_paragraph()
    pic.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_paragraph_format(pic, line_spacing=1.0, after_pt=6)
    pic.add_run().add_picture(BytesIO(png), width=Cm(width_cm))


def _para(
    doc: Document,
    text: str,
    style: dict[str, Any],
    *,
    bold: bool = False,
    italic: bool = False,
    center: bool = False,
    justify: bool = False,
    indent: bool = False,
    size: float | None = None,
    after: float | None = None,
    left_cm: float = 0,
    hanging_cm: float = 0,
) -> Any:
    para = doc.add_paragraph()
    if center:
        para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    elif justify:
        para.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    first = style["indent"] if indent else 0
    if hanging_cm:
        first = -hanging_cm
    set_paragraph_format(
        para,
        line_spacing=style["line"],
        after_pt=style["after"] if after is None else after,
        first_line_cm=first,
    )
    if left_cm:
        para.paragraph_format.left_indent = Cm(left_cm)
    _write_runs(para, text, style, bold=bold, italic=italic, size=size)
    return para


def _render_passage_blocks(doc: Document, blocks: list[dict[str, str]], style: dict[str, Any]) -> None:
    for block in blocks:
        kind = block.get("kind") or "body"
        if kind == "heading":
            _para(doc, _strip_markup(block.get("text") or ""), style, bold=True, after=4)
            continue
        if kind == "bullet":
            lead = (block.get("lead") or "").strip()
            text = (block.get("text") or "").strip()
            para = _para(doc, "", style, justify=True, after=2, left_cm=0.5, hanging_cm=0.37)
            _write_runs(para, "• ", style, bold=True)
            if lead:
                _write_runs(para, lead, style, bold=True)
            if text:
                _write_runs(para, (" " if lead else "") + text, style)
            continue
        indent = kind != "indent"
        headingish = bool(
            re.match(r"^\[\d+\]", block.get("text") or "")
            or re.match(r"^(A\d|N\d)\s*:", block.get("text") or "")
        )
        _para(
            doc,
            block.get("text") or "",
            style,
            justify=not headingish,
            indent=indent and not headingish,
            after=0,
        )


def _cm_dxa(cm: float) -> str:
    return str(int(round(cm * 567.0)))


def _set_col_widths(table, widths_cm: list[float]) -> None:
    table.autofit = False
    table.allow_autofit = False
    tbl = table._tbl
    tbl_pr = tbl.tblPr
    layout = tbl_pr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), _cm_dxa(sum(widths_cm)))
    tbl_w.set(qn("w:type"), "dxa")
    grid = tbl.find(qn("w:tblGrid"))
    if grid is not None:
        for child in list(grid):
            grid.remove(child)
        for width in widths_cm:
            col = OxmlElement("w:gridCol")
            col.set(qn("w:w"), _cm_dxa(width))
            grid.append(col)
    for row in table.rows:
        for cell, width in zip(row.cells, widths_cm):
            cell.width = Cm(width)
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), _cm_dxa(width))
            tc_w.set(qn("w:type"), "dxa")


def _set_table_borders(table, *, visible: bool = True, fill: str = "") -> None:
    tbl_pr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    outer = "single" if visible else "nil"
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        inner = edge.startswith("inside")
        el.set(qn("w:val"), "nil" if inner or not visible else outer)
        el.set(qn("w:sz"), "8" if el.get(qn("w:val")) == "single" else "0")
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), "000000" if el.get(qn("w:val")) == "single" else "auto")
        borders.append(el)
    tbl_pr.append(borders)
    if fill:
        shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear")
        shd.set(qn("w:color"), "auto")
        shd.set(qn("w:fill"), fill)
        tbl_pr.append(shd)


def _set_cell_margins(cell) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    mar = OxmlElement("w:tcMar")
    for edge, val in (("top", "60"), ("left", "80"), ("bottom", "60"), ("right", "80")):
        node = OxmlElement(f"w:{edge}")
        node.set(qn("w:w"), val)
        node.set(qn("w:type"), "dxa")
        mar.append(node)
    tc_pr.append(mar)


def _set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)
    _set_cell_margins(cell)


def _set_cell_borders(cell, *, color: str = "000000", fill: str = "") -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    borders = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "8")
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), color)
        borders.append(el)
    tc_pr.append(borders)
    if fill:
        _set_cell_shading(cell, fill)
    else:
        _set_cell_margins(cell)


def _fill_cell(cell, text: str, style: dict[str, Any], *, italic: bool = False, center: bool = False) -> None:
    cell.text = ""
    para = cell.paragraphs[0]
    if center:
        para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_paragraph_format(para, line_spacing=style["line"], after_pt=0)
    _write_runs(para, text, style, italic=italic)


def _cant_split_row(row) -> None:
    tr = row._tr
    tr_pr = tr.get_or_add_trPr()
    node = OxmlElement("w:cantSplit")
    tr_pr.append(node)


def _add_bordered_lines(doc: Document, lines: list[str], style: dict[str, Any]) -> None:
    table = doc.add_table(rows=1, cols=1)
    _set_col_widths(table, [CONTENT_WIDTH_CM])
    _set_table_borders(table, visible=True, fill="F3F3F3")
    cell = table.cell(0, 0)
    _set_cell_borders(cell)
    _cant_split_row(table.rows[0])
    cell.text = ""
    first = cell.paragraphs[0]
    set_paragraph_format(first, line_spacing=style["line"], after_pt=2)
    first.paragraph_format.keep_together = True
    _write_runs(first, lines[0], style)
    for line in lines[1:]:
        para = cell.add_paragraph()
        set_paragraph_format(para, line_spacing=style["line"], after_pt=2)
        para.paragraph_format.keep_together = True
        _write_runs(para, line, style)
    spacer = doc.add_paragraph()
    set_paragraph_format(spacer, line_spacing=style["line"], after_pt=4)


def _is_choice_line(text: str) -> bool:
    return bool(CHOICE_RE.match(_strip_markup(text).strip()))


WORD_BANK_STOP = {
    "a", "an", "the", "of", "to", "in", "on", "for", "and", "or", "but",
    "was", "were", "is", "are", "be", "been", "that", "this", "with",
    "from", "about", "into", "than", "then", "as", "at", "by", "it",
    "he", "she", "they", "we", "you", "his", "her", "their", "not",
}


def _is_word_bank_line(text: str) -> bool:
    stripped = _strip_markup(text.strip())
    if not stripped or _is_circled_item(stripped) or NUMBERED_RE.match(stripped) or CHOICE_RE.match(stripped):
        return False
    if "______" in stripped or "___" in stripped:
        return False
    words = stripped.split()
    if not (2 <= len(words) <= 6):
        return False
    if any(w in WORD_BANK_STOP for w in words):
        return False
    return all(re.fullmatch(r"[a-z][a-z'-]*", w) for w in words)


def _is_forms_wrap(text: str) -> bool:
    stripped = _strip_markup(text.strip())
    if not stripped or _is_circled_item(stripped) or NUMBERED_RE.match(stripped):
        return False
    if stripped.startswith("______") or stripped.startswith("___"):
        return True
    return bool(re.match(r"^_{3,}", stripped) and re.search(r"\b(adj|n|v|adv|prep)\.?", stripped))


def _add_word_bank_box(doc: Document, words: list[str], style: dict[str, Any]) -> None:
    _render_vocab_option_table(doc, words, style)


def _render_vocab_option_table(doc: Document, words: list[str], style: dict[str, Any]) -> None:
    cells = [str(w).strip() for w in words if str(w).strip()]
    if not cells:
        return
    if not _is_labeled_word_bank(cells):
        table = doc.add_table(rows=1, cols=1)
        _set_col_widths(table, [CONTENT_WIDTH_CM])
        _set_table_borders(table, visible=True)
        cell = table.cell(0, 0)
        _set_cell_borders(cell, color="000000")
        _fill_cell(cell, "      ".join(cells), style, italic=True, center=True)
        spacer = doc.add_paragraph()
        set_paragraph_format(spacer, line_spacing=style["line"], after_pt=4)
        return
    while len(cells) < 12:
        cells.append("")
    table = doc.add_table(rows=2, cols=6)
    col_w = CONTENT_WIDTH_CM / 6
    _set_col_widths(table, [col_w] * 6)
    _set_table_borders(table, visible=True)
    for row_i in range(2):
        for col_i in range(6):
            text = cells[row_i * 6 + col_i]
            cell = table.cell(row_i, col_i)
            _set_outer_box_cell(
                cell,
                top=row_i == 0,
                bottom=row_i == 1,
                left=col_i == 0,
                right=col_i == 5,
            )
            _fill_cell(cell, text, style)
    spacer = doc.add_paragraph()
    set_paragraph_format(spacer, line_spacing=style["line"], after_pt=6)


def _set_outer_box_cell(cell, *, top: bool, bottom: bool, left: bool, right: bool) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = OxmlElement("w:tcBorders")
    for edge, on in (("top", top), ("left", left), ("bottom", bottom), ("right", right)):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single" if on else "nil")
        el.set(qn("w:sz"), "8" if on else "0")
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), "000000" if on else "auto")
        borders.append(el)
    tc_pr.append(borders)
    _set_cell_margins(cell)


def _render_exercise_lines(doc: Document, lines: list[str], style: dict[str, Any]) -> None:
    boxed: list[str] = []
    choices: list[str] = []

    def flush_box() -> None:
        nonlocal boxed
        if boxed:
            _add_bordered_lines(doc, boxed, style)
            boxed = []

    def flush_choices() -> None:
        nonlocal choices
        if choices:
            _add_bordered_lines(doc, choices, style)
            choices = []

    for line in lines:
        text = str(line).strip()
        if not text or re.fullmatch(r"丂+", text.replace(" ", "")):
            continue
        plain = _strip_markup(text)
        if _is_circled_item(text) or (boxed and _is_forms_wrap(text)):
            flush_choices()
            boxed.append(text)
            continue
        flush_box()
        if _is_choice_line(text) or (
            choices
            and not NUMBERED_RE.match(plain)
            and not _is_circled_item(plain)
            and not SECTION_RE.match(plain)
        ):
            choices.append(text)
            continue
        flush_choices()
        if _is_word_bank_line(text):
            _add_word_bank_box(doc, _strip_markup(text).split(), style)
            continue
        _para(doc, text, style, justify=not text.startswith("_"), after=2)
    flush_box()
    flush_choices()


def _is_cloze_choice_set(questions: list[dict[str, Any]]) -> bool:
    if len(questions) < 8:
        return False
    for question in questions:
        if (question.get("stem") or "").strip():
            return False
        if len(question.get("options") or []) != 4:
            return False
    return True


def _render_cloze_choice_table(doc: Document, questions: list[dict[str, Any]], style: dict[str, Any]) -> None:
    num_w = 1.15
    opt_w = (CONTENT_WIDTH_CM - num_w) / 4
    widths = [num_w, opt_w, opt_w, opt_w, opt_w]
    table = doc.add_table(rows=len(questions), cols=5)
    _set_col_widths(table, widths)
    _set_table_borders(table, visible=False)
    for i, question in enumerate(questions):
        _fill_cell(table.cell(i, 0), f"{question.get('no')}.", style)
        options = list(question.get("options") or [])
        by_label = {str(o.get("label") or ""): str(o.get("text") or "").strip() for o in options}
        for j, label in enumerate("ABCD"):
            _fill_cell(table.cell(i, j + 1), f"{label}. {by_label.get(label, '')}".strip(), style)
    spacer = doc.add_paragraph()
    set_paragraph_format(spacer, line_spacing=style["line"], after_pt=4)


def _render_choices(doc: Document, options: list[dict[str, str]], style: dict[str, Any]) -> None:
    texts = [f"{o.get('label', '')}. {o.get('text', '')}".strip() for o in options]
    if not texts:
        return
    longest = max(len(t) for t in texts)
    if 2 <= len(texts) <= 4 and longest <= 28:
        widths = [CONTENT_WIDTH_CM / len(texts)] * len(texts)
        table = doc.add_table(rows=1, cols=len(texts))
        _set_col_widths(table, widths)
        _set_table_borders(table, visible=False)
        for i, text in enumerate(texts):
            _fill_cell(table.cell(0, i), text, style)
        spacer = doc.add_paragraph()
        set_paragraph_format(spacer, line_spacing=style["line"], after_pt=2)
        return
    if len(texts) == 4 and longest <= 55:
        table = doc.add_table(rows=2, cols=2)
        _set_col_widths(table, [CONTENT_WIDTH_CM / 2, CONTENT_WIDTH_CM / 2])
        _set_table_borders(table, visible=False)
        cells = [texts[0], texts[1], texts[2], texts[3]]
        _fill_cell(table.cell(0, 0), cells[0], style)
        _fill_cell(table.cell(0, 1), cells[1], style)
        _fill_cell(table.cell(1, 0), cells[2], style)
        _fill_cell(table.cell(1, 1), cells[3], style)
        spacer = doc.add_paragraph()
        set_paragraph_format(spacer, line_spacing=style["line"], after_pt=4)
        return
    for text in texts:
        _para(doc, text, style, after=0)
    spacer = doc.add_paragraph()
    set_paragraph_format(spacer, line_spacing=style["line"], after_pt=4)
