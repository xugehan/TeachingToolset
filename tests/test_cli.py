from pathlib import Path

import pytest
from docx import Document

from english_toolkit.cli import main
from english_toolkit.dictation import make_dictation_pdf
from english_toolkit.exam import EXAMPLE_EXAM, export_exam_docx
from english_toolkit.grades import generate_grades_pdf, read_grades_excel


def test_dictation_pdf(tmp_path: Path) -> None:
    out = tmp_path / "dictation.pdf"
    make_dictation_pdf(
        "0831",
        "eager-effort",
        ["n. 鹰", "赢得好名声", "adj. 高效的"],
        str(out),
        cols=2,
        rows=3,
    )
    data = out.read_bytes()
    assert data.startswith(b"%PDF")
    assert len(data) > 1000


def test_grades_pdf_from_template(tmp_path: Path) -> None:
    xlsx = tmp_path / "grades.xlsx"
    pdf = tmp_path / "grades.pdf"
    assert main(["template", "grades", "-o", str(xlsx)]) == 0
    df = read_grades_excel(str(xlsx))
    generate_grades_pdf(df, str(pdf), card_title="期中英语")
    assert pdf.read_bytes().startswith(b"%PDF")


def test_exam_docx_and_answers(tmp_path: Path) -> None:
    paper = tmp_path / "exam.docx"
    answers = tmp_path / "answers.docx"
    export_exam_docx(EXAMPLE_EXAM, str(paper), answers=False)
    export_exam_docx(EXAMPLE_EXAM, str(answers), answers=True)
    paper_doc = Document(str(paper))
    answer_doc = Document(str(answers))
    paper_text = "\n".join(p.text for p in paper_doc.paragraphs)
    answer_text = "\n".join(p.text for p in answer_doc.paragraphs)
    assert "Listening Comprehension" in paper_text
    assert "Translation" in paper_text
    assert "只要你坚持练习" in paper_text
    assert "答案" in answer_text
    assert "听力原文" in answer_text


def test_cli_exam_template_and_export(tmp_path: Path) -> None:
    json_path = tmp_path / "exam.json"
    docx_path = tmp_path / "out.docx"
    assert main(["template", "exam", "-o", str(json_path)]) == 0
    assert main(["exam", "-i", str(json_path), "-o", str(docx_path), "--answers"]) == 0
    assert docx_path.exists()
    assert (tmp_path / "out-答案.docx").exists()


def test_cli_dictation_from_text(tmp_path: Path) -> None:
    items = tmp_path / "items.txt"
    items.write_text("n. 鹰\n# skip\n赢得好名声\n", encoding="utf-8")
    out = tmp_path / "d.pdf"
    assert main([
        "dictation",
        "-i",
        str(items),
        "--date",
        "0831",
        "--scope",
        "Unit1",
        "-o",
        str(out),
    ]) == 0
    assert out.exists()


def test_exam_catalog_payload() -> None:
    from english_toolkit.catalog import build_exam_payload

    data = build_exam_payload(
        title="期中考试",
        total_score=150,
        duration="120分钟",
        year_month="2026.11",
        selected_ids=["listening_a", "translation"],
        contents={"listening_a": "1. A. Hello.", "translation": "1. 坚持练习。"},
        instructions={"listening_a": "Directions: ...", "translation": "Translate."},
        answers={"listening_a": "1. A", "translation": "Keep practising."},
        listening_scripts="W: Hello.",
    )
    assert data["exam"]["title"] == "期中考试"
    assert data["sections"][0]["major"] == "I"
    assert data["sections"][1]["major"] == "II"
    assert data["sections"][1]["type"] == "translation"
    assert data["listening_scripts"][0]["text"] == "W: Hello."


def test_exam_catalog_requires_selection() -> None:
    from english_toolkit.catalog import build_exam_payload

    with pytest.raises(ValueError):
        build_exam_payload(
            title="x",
            total_score=100,
            duration="",
            year_month="",
            selected_ids=[],
            contents={},
            instructions={},
        )
