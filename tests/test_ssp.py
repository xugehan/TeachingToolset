from pathlib import Path
import re

import pytest
from docx import Document

from english_toolkit.mkyice import _error_message, bubble_items, choice_bits
from english_toolkit.ssp import _strip_markup, collect_mcqs, describe_exercises, extract_issue, export_reading_docx
from english_toolkit.ssp_answers import parse_answer_text

SAMPLE_PDFS = [
    Path("/Users/claire/Desktop/高二上/试卷/平时作业/0911周末/1788832552588_7e531b42c3913b24.pdf"),
    Path(
        "/Users/claire/.cursor/projects/Users-claire-code-english-toolkit-cli/"
        "attachments/empty-state-draft/1788832552588_7e531b42c3913b24.pdf"
    ),
]


def _sample_pdf() -> Path:
    for path in SAMPLE_PDFS:
        if path.exists():
            return path
    pytest.skip("SSP 第1647期 PDF 不在常见路径")


def test_extract_issue_attaches_b_page_questions() -> None:
    issue = extract_issue(_sample_pdf())
    assert issue["issue"] == "1647"
    by_title = {article["title"]: article for article in issue["articles"]}

    joy = by_title["Can Micro-Acts of Joy Make You Happier?"]
    assert "4道选择" in describe_exercises(joy)
    joy_mcq = [q for block in joy["exercises"] for q in block.get("questions") or [] if q.get("options")]
    assert len(joy_mcq) == 4
    assert joy_mcq[0]["options"][0]["label"] == "A"

    mom = next(article for title, article in by_title.items() if title.startswith("My Mom"))
    assert "Guide" in mom["title"]
    box = "\n".join(line for block in mom["exercises"] for line in block.get("lines") or [])
    assert "A. Therefore" in box
    assert "F. And this kind of savoring" in box

    heat = by_title["How Heat Domes Are Reshaping Summer"]
    guess = heat["exercises"][0]["questions"]
    assert len(guess[0]["options"]) == 3
    assert heat["diagrams"]
    heat_mcq = heat["exercises"][-1]["questions"]
    assert len(heat_mcq) == 4
    assert "diagram" in heat_mcq[0]["stem"].lower()


def test_export_reading_docx_includes_passages_and_questions(tmp_path: Path) -> None:
    issue = extract_issue(_sample_pdf())
    wanted = [
        "Can Micro-Acts of Joy Make You Happier?",
        "My Mom's Guide to the Art of Living",
        "How Heat Domes Are Reshaping Summer",
    ]
    ids = []
    for title in wanted:
        article = next(item for item in issue["articles"] if item["title"].startswith(title[:12]))
        ids.append(article["id"])
    out = tmp_path / "ssp.docx"
    export_reading_docx(issue, ids, out, include_questions=True, include_vocab_exercises=True)
    doc = Document(str(out))
    text = "\n".join(para.text for para in doc.paragraphs)
    assert "Passage 1" in text
    assert "Passage 2" in text
    assert "Passage 3" in text
    assert "Directions" not in text
    assert "III. Choose the best answer." in text
    assert "___1___" in text
    assert doc.inline_shapes


SAMPLE_PDFS_1648 = [
    Path("/Users/claire/Desktop/高二上/试卷/平时作业/0918周末/1789378407123_63a4a63cdc497f82.pdf"),
]


def _sample_pdf_1648() -> Path:
    for path in SAMPLE_PDFS_1648:
        if path.exists():
            return path
    pytest.skip("SSP 第1648期 PDF 不在常见路径")


def test_extract_volunteer_layout_and_cloze() -> None:
    issue = extract_issue(_sample_pdf_1648())
    assert issue["issue"] == "1648"
    volunteer = next(a for a in issue["articles"] if "Animal Services" in a["title"])
    kinds = [b.get("kind") for b in volunteer["passage_blocks"]]
    assert "heading" in kinds
    assert "bullet" in kinds
    assert any(b.get("text") == "What do volunteers do?" for b in volunteer["passage_blocks"])
    move = next(b for b in volunteer["passage_blocks"] if b.get("lead") == "Like to move?")
    assert "狗窝" in (move.get("text") or "")
    circled = "\n".join(ln for block in volunteer["exercises"] for ln in block.get("lines") or [])
    assert "① confusing" in circled
    assert "② scary" in circled

    cloze = next(a for a in issue["articles"] if a.get("kind") == "extra")
    assert cloze["title"] == "Winning Life's Gold Medal"
    volunteering = next(a for a in issue["articles"] if a["title"].startswith("Want Real-Life"))
    bodies = [b for b in volunteering["passage_blocks"] if b.get("kind") == "body"]
    assert len(bodies) >= 8
    assert bodies[0]["text"].startswith("For the last year")
    guess = [q for block in volunteering["exercises"] for q in block.get("questions") or [] if q.get("options")]
    stem1 = next(q["stem"] for q in guess if str(q.get("no")) == "1")
    stem2 = next(q["stem"] for q in guess if str(q.get("no")) == "2")
    assert "⟪gotten her hands dirty⟫" in stem1
    assert "⟪uptick⟫" in stem2
    assert any("However, challenges" in (b.get("text") or "") for b in bodies)
    blob = "\n".join(_strip_markup(b.get("text") or "") for b in bodies)
    assert '"You\'re not doomscrolling' in blob
    assert not any(ch in blob for ch in "‘’“”")
    assert any("To attract young people" in (b.get("text") or "") for b in bodies)
    assert cloze["genre"] == "完形填空"
    assert "15道选择" in describe_exercises(cloze)
    questions = [q for block in cloze["exercises"] for q in block.get("questions") or []]
    assert len(questions) == 15
    assert [o["label"] for o in questions[0]["options"]] == ["A", "B", "C", "D"]
    assert questions[0]["options"][0]["text"] == "appearance"


def test_export_volunteer_boxes_and_cloze(tmp_path: Path) -> None:
    issue = extract_issue(_sample_pdf_1648())
    volunteer = next(a for a in issue["articles"] if "Animal Services" in a["title"])
    cloze = next(a for a in issue["articles"] if a.get("kind") == "extra")
    out = tmp_path / "ssp1648.docx"
    export_reading_docx(issue, [volunteer["id"], cloze["id"]], out)
    doc = Document(str(out))
    text = "\n".join(para.text for para in doc.paragraphs)
    assert "What do volunteers do?" in text
    assert "• Like to move?" in text
    assert "kennels (狗窝)" in text
    assert "Winning Life's Gold Medal" in text
    assert "Cloze" in text
    boxed = doc.tables[0].cell(0, 0).text
    assert "① confusing" in boxed
    assert "② scary" in boxed
    assert "1. The sudden loud noise" not in boxed
    bank = next(
        table.cell(0, 0).text
        for table in doc.tables
        if "comfort" in table.cell(0, 0).text and "potential" in table.cell(0, 0).text
    )
    assert "update" in bank
    cloze_tbl = max(doc.tables, key=lambda table: len(table.rows))
    assert len(cloze_tbl.rows) == 15
    assert len(cloze_tbl.columns) == 5
    assert cloze_tbl.cell(0, 0).text.strip() == "4."
    assert cloze_tbl.cell(0, 1).text.startswith("A. appearance")
    assert cloze_tbl.cell(0, 2).text.startswith("B. performance")
    assert cloze_tbl.cell(0, 3).text.startswith("C. quality")
    assert cloze_tbl.cell(0, 4).text.startswith("D. status")
    assert cloze_tbl.cell(1, 1).text.startswith("A. action")


def test_extract_later_articles_keep_newspaper_paragraphs() -> None:
    issue = extract_issue(_sample_pdf_1648())
    zendaya = next(a for a in issue["articles"] if "Zendaya" in a["title"])
    passage = zendaya["passage"]
    assert "Spider-Man" in passage
    assert "pizza" not in passage.lower()
    assert len([b for b in zendaya["passage_blocks"] if b.get("kind") == "body"]) >= 6

    sky = next(a for a in issue["articles"] if "Sky" in a["title"])
    bodies = [_strip_markup(b.get("text") or "") for b in sky["passage_blocks"] if b.get("kind") == "body"]
    assert bodies[0].startswith("⑤ If the architects")
    assert "orbit (轨道)" in bodies[0]
    assert "The main benefit" in bodies[1]
    assert not any("driven past a big windowless" in (b or "") for b in bodies)
    bank = [
        _strip_markup(ln)
        for block in sky["exercises"]
        for ln in block.get("lines") or []
        if re.match(r"^[A-F]\.", _strip_markup(ln).strip())
    ]
    assert len(bank) == 6
    assert bank[0].startswith("A. Yet earthbound")
    assert "opportunity" in bank[3]
    assert bank[3].startswith("D. ")
    assert not any(ln.startswith("opportunity") for ln in bank)

    friendship = next(a for a in issue["articles"] if "Friendship" in a["title"])
    assert _strip_markup(friendship["passage_blocks"][0]["text"]).startswith("On the surface")
    assert "Ono fancy" not in friendship["passage"]


def test_export_guess_meaning_is_bold_underline(tmp_path: Path) -> None:
    issue = extract_issue(_sample_pdf_1648())
    volunteering = next(a for a in issue["articles"] if a["title"].startswith("Want Real-Life"))
    out = tmp_path / "ssp-underline.docx"
    export_reading_docx(issue, [volunteering["id"]], out)
    doc = Document(str(out))
    marked = []
    for para in doc.paragraphs:
        for run in para.runs:
            if run.underline and run.bold and run.text.strip():
                marked.append(run.text)
    joined = " ".join(marked)
    assert "gotten her hands dirty" in joined
    assert "uptick" in joined


def test_export_instructions_are_bold_not_italic(tmp_path: Path) -> None:
    issue = extract_issue(_sample_pdf_1648())
    volunteer = next(a for a in issue["articles"] if "Animal Services" in a["title"])
    out = tmp_path / "ssp-instructions.docx"
    export_reading_docx(issue, [volunteer["id"]], out)
    doc = Document(str(out))
    titles = [p for p in doc.paragraphs if p.text.startswith("I.") or p.text.startswith("II.")]
    assert titles
    for para in titles:
        assert any(run.bold for run in para.runs)
        assert not any(run.italic for run in para.runs if run.text.strip())


def test_export_numbers_mcqs_sequentially_and_leaves_fillins(tmp_path: Path) -> None:
    issue = extract_issue(_sample_pdf_1648())
    volunteering = next(a for a in issue["articles"] if a["title"].startswith("Want Real-Life"))
    animal = next(a for a in issue["articles"] if "Animal Services" in a["title"])
    out = tmp_path / "ssp-seq.docx"
    export_reading_docx(issue, [volunteering["id"], animal["id"]], out)
    doc = Document(str(out))
    text = "\n".join(para.text for para in doc.paragraphs)
    assert "1. She has learned about birds" in text
    assert "3. What did Breece Eagar gain" in text
    assert "7. Which of the following is true about the volunteer work" in text
    assert "1. The sudden loud noise" in text
    assert "1. What did Breece Eagar gain" not in text


def test_parse_answer_text_assigns_section_keys() -> None:
    parsed = parse_answer_text(
        """
Become a Volunteer for Animal Services (A4)
I.
1. scared
III.
1-3 ADD
Want Real-Life Connection? Try Volunteering (A5)
I.
1-2 BB
III.
1-4 DDAB
Sentence Structures
1. He was in the midst
Cloze
Winning Life's Gold Medal
1-5 BBADD
6-10 CBCBA
11-15 BABCA
"""
    )
    by_title = {a["title"]: a for a in parsed["articles"]}
    assert by_title["Become a Volunteer for Animal Services"]["sections"]["III"]["mcq"] == "ADD"
    assert by_title["Want Real-Life Connection? Try Volunteering"]["sections"]["I"]["mcq"] == "BB"
    assert by_title["Want Real-Life Connection? Try Volunteering"]["sections"]["III"]["mcq"] == "DDAB"
    assert by_title["Winning Life's Gold Medal"]["sections"]["Cloze"]["mcq"] == "BBADDCBCBABABCA"


def test_choice_bits_marks_single_answer() -> None:
    assert choice_bits("B", 4) == "0100"
    assert choice_bits("C", 3) == "001"
    assert choice_bits("E", 6) == "000010"
    assert choice_bits("", 4) == "0000"


def test_bubble_items_keeps_guess_three_and_six_choose_six() -> None:
    items = bubble_items(
        [
            {"no": 4, "options": [{"label": "A"}, {"label": "B"}, {"label": "C"}], "answer": "B"},
            {"no": 6, "options": [{"label": "A"}, {"label": "B"}, {"label": "C"}, {"label": "D"}], "answer": "D"},
            {
                "no": 14,
                "kind": "six",
                "options": [{"label": label} for label in "ABCDEF"],
                "answer": "E",
            },
        ]
    )
    assert items[0]["choices"] == "010"
    assert items[1]["choices"] == "0001"
    assert items[2]["choices"] == "000010"
    assert items[2]["itemOrdinal"] == 14


def test_mkyice_error_message_maps_password_errors() -> None:
    assert "密码不对" in _error_message(
        422, '{"ERRORS":[{"ERROR_TYPE":"SESSION___LOGIN___PASSWORD_NOTCORRECT"}]}'
    )
    assert "会话" in _error_message(400, '{"ERRORS":["invalid teacherUserId"]}')


def test_six_choose_is_last_reading_mcq_with_six_options() -> None:
    issue = extract_issue(_sample_pdf_1648())
    sky = next(a for a in issue["articles"] if "Sky" in a["title"])
    gaudi = next(a for a in issue["articles"] if "Gaudí" in a["title"] or "Gaudi" in a["title"])
    cloze = next(a for a in issue["articles"] if a.get("kind") == "extra")
    mcqs = collect_mcqs(issue, [sky["id"], gaudi["id"], cloze["id"]])
    reading = [q for q in mcqs if not q.get("is_cloze")]
    six = [q for q in reading if q.get("kind") == "six"]
    assert len(six) == 4
    assert reading[-4:] == six
    assert [o["label"] for o in six[0]["options"]] == ["A", "B", "C", "D", "E", "F"]
    assert all(q.get("is_cloze") for q in mcqs[-15:])


def test_sports_article_keeps_full_two_column_passage() -> None:
    issue = extract_issue(_sample_pdf())
    sports = next(a for a in issue["articles"] if a["title"].startswith("Why Watching Sports"))
    text = sports["passage"]
    assert "Keyes is one of many psychologists" in text
    assert "rituals are comforting" in text
    assert "net benefit" in text
    assert len(text) > 1500
    assert not text.startswith("A. applied")


def test_banked_cloze_uses_labeled_word_table(tmp_path: Path) -> None:
    issue = extract_issue(_sample_pdf())
    extra = next(
        a
        for a in issue["articles"]
        if a.get("kind") == "extra" and "Great White" in a.get("title", "")
    )
    assert extra["genre"] == "选词填空"
    assert extra["word_bank"][0].startswith("A. ")
    assert extra["word_bank"][10].startswith("K. ")
    assert len(extra["word_bank"]) == 11
    assert extra["passage"].startswith("My wife, Kay")
    assert "A. applied" not in extra["passage"]
    out = tmp_path / "vocab.docx"
    export_reading_docx(issue, [extra["id"]], out, include_questions=True)
    doc = Document(str(out))
    text = "\n".join(para.text for para in doc.paragraphs)
    assert "Meeting the Great White" in text
    assert "My wife, Kay" in text
    assert doc.tables
    bank = " ".join(cell.text.strip() for row in doc.tables[0].rows for cell in row.cells)
    assert "A. applied" in bank
    assert "K. wandered" in bank
